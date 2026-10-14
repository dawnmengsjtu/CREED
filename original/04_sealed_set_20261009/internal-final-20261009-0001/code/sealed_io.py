"""Identity-gated loaders for the sealed internal-final one-shot (PLAN.md §2).

Rules enforced here (hard):
  * Every raw-file reader receives an explicit `allowed` ID set; assay fields of any other line are never converted to numbers,
    never stored and never printed (row / column gating identical to generator-fullpool-20261009-0001/code/extract_fullpool.py and to
    the original readers in creed/science/ccle_values.py).
  * The caller passes the forbidden set (the 64 validation IDs); an `allowed` set that intersects it is refused before any file opens.
  * The AUC reader is a verbatim copy of original_prepare_gdsc421.py::extract; the CRISPR reader mirrors
    prepare_crispr_development.py (index_models + extract). Both expose parity checks that run on the 421 / 310 development
    lines only.
"""
import csv
import gzip
import hashlib
import json
import os
import re
from collections import Counter, defaultdict

import numpy as np

RAW_SHA = {'CCLE_metabolomics_20190502.csv': '7c1d24aa575f4c58a29019026b5df8e6d1142a56925aba32ff3f1d1d5a7fd0ac',
           'CCLE_RNAseq_RSEM_Expected_Counts.csv.gz': 'f59221c04e15a2ba807181d23ba53fa4acd778f02ecb05321c473857e824191f',
           'mmc2.xlsx': '9327f665d259ef4d148d3a64e7c76cdeb51dce8dee06768efd6b7d044130e203',
           'mmc3.xlsx': 'f6c414bc07fdf6e33680cd61c4b3d22d26e531a27eff2cdef3648bfb46e72f34',
           'Model.csv': 'a4cac376131b41aa10b60a075b11c80264bfa860a5509d22cde259c5e85867f8'}
GDSC_24JUL22_SHA = '9d6dc82b9a072bb00b6139dcd6c4c11de1e047f210382fb726e92ab8dd0eae2a'
DEPMAP_24Q2_EFFECT_MD5 = 'd10d73692672380b3792231549e97d5a'
DEPMAP_24Q2_EFFECT_BYTES = 419330454
PARITY = ['SIDM00018', 'SIDM00080', 'SIDM00082', 'SIDM00396', 'SIDM00685']   # 421-train lines (generator-fullpool parity set)
M_FEATURES = 225
RNA_SOURCE_GENES = 37279


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def md5(path):
    h = hashlib.md5()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def prefix_fields(raw, last):
    """Quote-aware byte split of the first `last`+1 CSV fields (verbatim extract_fullpool.py)."""
    out = []; start = 0; i = 0; quoted = False
    while i < len(raw):
        c = raw[i]
        if c == 34:
            if quoted and i + 1 < len(raw) and raw[i + 1] == 34:
                i += 2; continue
            quoted = not quoted
        if c == 44 and not quoted:
            out.append(raw[start:i]); start = i + 1
            if len(out) > last:
                return out
        i += 1
    out.append(raw[start:].rstrip(b'\r\n')); assert len(out) > last and not quoted
    return out


def text(v):
    return next(csv.reader([v.decode('utf-8-sig')]))[0].strip()


def number(v):
    if v.strip().lower() in ('', 'na', 'nan', 'null', 'none'):
        return np.nan
    return float(v)


class Gate(Exception):
    pass


def check_allowed(allowed, forbidden, universe=None):
    allowed = list(allowed)
    if len(allowed) != len(set(allowed)) or not allowed:
        raise Gate('allowed ID list must be non-empty and unique')
    if any(not re.fullmatch(r'SIDM\d{5}', s) for s in allowed):
        raise Gate('allowed IDs must be SIDM identifiers')
    bad = sorted(set(allowed) & set(forbidden))
    if bad:
        raise Gate(f'refusing: {len(bad)} allowed IDs are in the forbidden (validation) set')
    if universe is not None:
        extra = sorted(set(allowed) - set(universe))
        if extra:
            raise Gate(f'refusing: {len(extra)} allowed IDs are outside the declared universe (train ∪ internal_final)')
    return set(allowed)


# ---------------------------------------------------------------- identity tables (IDs + metadata only)
def read_mmc2(path):
    """mmc2 'Cell line level sample info': SIDM -> {cell_line, tissue, achs, ccles}. No assay values exist in this sheet."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb['Cell line level sample info']
    rows = ws.iter_rows(values_only=True)
    next(rows); hdr = [str(c) if c is not None else '' for c in next(rows)]
    col = {k: hdr.index(k) for k in ('model_id', 'Cell_line', 'Tissue_type', 'BROAD_ID', 'CCLE_ID')}
    mmc2 = {}
    for r in rows:
        if r is None or r[col['model_id']] is None:
            continue
        sid = str(r[col['model_id']]).strip()
        assert re.fullmatch(r'SIDM\d{5}', sid) and sid not in mmc2, sid
        g = lambda k: ('' if r[col[k]] is None else str(r[col[k]]).strip())
        mmc2[sid] = dict(cell_line=g('Cell_line'), tissue=g('Tissue_type'), ach=g('BROAD_ID'), ccle=g('CCLE_ID'))
    wb.close()
    assert len(mmc2) == 949, len(mmc2)
    ach_split = defaultdict(set); ccle_split = defaultdict(set); ach_whole = {}; ccle_whole = {}
    for sid, r in mmc2.items():
        for a in [x.strip() for x in r['ach'].split(';') if x.strip()]:
            ach_split[a].add(sid)
        for c in [x.strip() for x in r['ccle'].split(';') if x.strip()]:
            ccle_split[c].add(sid)
        if r['ach'] and ';' not in r['ach']:
            assert r['ach'] not in ach_whole; ach_whole[r['ach']] = sid
        if r['ccle'] and ';' not in r['ccle']:
            assert r['ccle'] not in ccle_whole; ccle_whole[r['ccle']] = sid
    maps = dict(ach_split={a: next(iter(s)) for a, s in ach_split.items() if len(s) == 1},
                ccle_split={c: next(iter(s)) for c, s in ccle_split.items() if len(s) == 1},
                ach_whole=ach_whole, ccle_whole=ccle_whole)
    return mmc2, maps


def resolve_m_row(ccle, dep, maps):
    """mmc2-only resolution of one metabolomics row under BOTH conventions (split / whole-string); conflicts fail closed."""
    res = []
    for conv in ('split', 'whole'):
        a = maps['ach_' + conv].get(dep) if dep else None
        c = maps['ccle_' + conv].get(ccle) if ccle else None
        if a and c and a != c:
            res.append(None)
        else:
            res.append(a or c)
    if res[0] != res[1]:
        return None
    return res[0]


# ---------------------------------------------------------------- gated omics loader
def load_omics(raw, allowed, forbidden, universe, log=print):
    """M (225), RNA log-CPM (34,985 frozen common genes), P (6,692) and mmc2 tissue for exactly the `allowed` IDs.
    `raw` = dict with paths: metabolomics, rna, mmc2, mmc3, plus `common` (34,985 source gene indices) and `all_gene_ids`,
    `metabolite_ids`, `protein_ids` (frozen feature identities from the 421 artifacts)."""
    allowed = check_allowed(allowed, forbidden, universe)
    order = sorted(allowed)
    for name, path in (('CCLE_metabolomics_20190502.csv', raw['metabolomics']), ('CCLE_RNAseq_RSEM_Expected_Counts.csv.gz', raw['rna']),
                       ('mmc2.xlsx', raw['mmc2']), ('mmc3.xlsx', raw['mmc3'])):
        h = sha(path)
        if h != RAW_SHA[name]:
            raise Gate(f'raw file sha mismatch: {name} {h}')
    mmc2, maps = read_mmc2(raw['mmc2'])
    missing = [s for s in order if s not in mmc2]
    if missing:
        raise Gate(f'{len(missing)} allowed IDs are not in mmc2')
    tissues = np.array([mmc2[s]['tissue'] for s in order])
    idx = {s: i for i, s in enumerate(order)}
    # ---- metabolomics: identity pass, then value pass on allowed rows only
    mp = raw['metabolomics']
    with open(mp, 'rb') as f:
        header = next(csv.reader([f.readline().decode('utf-8-sig')]))
        if header[:2] != ['CCLE_ID', 'DepMap_ID'] or header[2:] != list(raw['metabolite_ids']):
            raise Gate('metabolomics header differs from the frozen 225 metabolite identities')
        rows = []
        for line in f:
            if not line.strip():
                continue
            p = prefix_fields(line, 1); rows.append((text(p[0]), text(p[1])))
    row2sid = {}
    for i, (ccle, dep) in enumerate(rows):
        s = resolve_m_row(ccle, dep, maps)
        if s in allowed:
            if s in row2sid.values():
                raise Gate(f'duplicate metabolomics row for an allowed line')
            row2sid[i] = s
    if set(row2sid.values()) != allowed:
        raise Gate(f'metabolomics rows resolved for {len(set(row2sid.values()))} of {len(allowed)} allowed IDs')
    m = np.full((len(order), M_FEATURES), np.nan)
    with open(mp, 'rb') as f:
        f.readline(); i = -1
        for line in f:
            if not line.strip():
                continue
            i += 1
            s = row2sid.get(i)
            if s is None:
                continue        # identity-only traversal: assay fields of every other line stay undecoded
            v = next(csv.reader([line.decode('utf-8-sig')])); assert len(v) == M_FEATURES + 2
            m[idx[s]] = [number(z) for z in v[2:]]
    assert not np.isinf(m).any()
    log(event='sealed_loader_M', rows=len(rows), decoded=len(row2sid))
    # ---- RNA: header identity, then column-gated decoding
    rp = raw['rna']
    with gzip.open(rp, 'rb') as f:
        head = [next(csv.reader([f.readline().decode('utf-8-sig')])) for _ in range(4)]
    if [h[0] for h in head] != ['model_id', 'model_name', 'dataset_name', 'gene_id']:
        raise Gate('unexpected RNA metadata rows')
    rna_cols = {s: j for j, s in enumerate(head[0][2:], 2)}
    missing = [s for s in order if s not in rna_cols]
    if missing:
        raise Gate(f'{len(missing)} allowed IDs have no RNA column')
    cols = [rna_cols[s] for s in order]; width = len(head[0])
    rna_batch = np.array([head[2][rna_cols[s]] for s in order])
    common = np.asarray(raw['common']); slot = {int(g): j for j, g in enumerate(common)}
    counts = np.full((len(order), len(common)), np.nan); genes = []
    with gzip.open(rp, 'rb') as f:
        for _ in range(4):
            f.readline()
        for line in f:
            if not line.strip():
                continue
            flds = line.rstrip(b'\r\n').split(b',')
            if len(flds) != width:
                flds = prefix_fields(line, width - 1)
            assert len(flds) == width
            gi = len(genes); genes.append(text(flds[0]) if flds[0].strip() else '')
            if gi in slot:
                vals = np.array([number(text(flds[j]) if flds[j].strip() else '') for j in cols], float)
                assert not np.isinf(vals).any() and not np.any(vals < 0)
                counts[:, slot[gi]] = vals
    if len(genes) != RNA_SOURCE_GENES or not np.array_equal(np.array(genes), np.asarray(raw['all_gene_ids'])):
        raise Gate('RNA source gene identities differ from the frozen 37,279 list')
    good = np.isfinite(counts).all(1) & (np.nansum(counts, 1) > 0)
    if not good.all():
        raise Gate(f'{int((~good).sum())} allowed lines have incomplete RNA on the frozen panel')
    rna = np.log1p(counts / counts.sum(1)[:, None] * 1e6)
    log(event='sealed_loader_RNA', columns=len(cols), genes=len(common))
    # ---- protein: mmc3 row-2 header identity, rows gated by SIDM prefix
    import openpyxl
    wb = openpyxl.load_workbook(raw['mmc3'], read_only=True, data_only=True)
    try:
        ws = wb['Prot matrix excl single-peptide']
        hdr = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        prot = tuple('' if v is None else str(v).strip() for v in hdr[1:])
        while prot and not prot[-1]:
            prot = prot[:-1]
        if str(hdr[0]).strip() != 'Project_Identifier' or list(prot) != list(raw['protein_ids']):
            raise Gate('mmc3 protein header differs from the frozen 6,692 protein identities')
        p = np.full((len(order), len(prot)), np.nan); seen = set()
        for row in ws.iter_rows(min_row=3, max_col=len(prot) + 1, values_only=True):
            if not row or row[0] is None:
                continue
            sid = str(row[0]).strip().split(';', 1)[0]
            if sid not in allowed:
                continue
            if sid in seen:
                raise Gate('duplicate mmc3 row for an allowed line')
            seen.add(sid)
            p[idx[sid]] = [np.nan if (x is None or (isinstance(x, str) and x.strip().lower() in ('', 'na', 'nan'))) else float(x) for x in row[1:]]
    finally:
        wb.close()
    if seen != allowed:
        raise Gate(f'mmc3 rows found for {len(seen)} of {len(allowed)} allowed IDs')
    assert not np.isinf(p).any()
    log(event='sealed_loader_P', rows=len(seen), proteins=len(prot))
    return dict(sample_ids=np.array(order), m=m, rna=rna, p=p, tissues=tissues, rna_batch=rna_batch)


def omics_parity(D, S, R421):
    """Compare the PARITY lines of a loaded omics dict with the 421 artifacts (M: source-inputs, RNA: rna421, P: source-inputs)."""
    si = {s: i for i, s in enumerate(S['sample_ids'].tolist())}; ri = {s: i for i, s in enumerate(R421['sample_ids'].tolist())}
    di = {s: i for i, s in enumerate(D['sample_ids'].tolist())}
    ids = [s for s in PARITY if s in di]
    out = dict(parity_ids=ids, m_max_abs_diff=0., m_nan_equal=True, rna_max_abs_diff=0., p_max_abs_diff=0., p_nan_equal=True, tissue_equal=True)
    for s in ids:
        a, b = D['m'][di[s]], S['m'][si[s]]
        out['m_nan_equal'] &= bool(np.array_equal(np.isnan(a), np.isnan(b))); out['m_max_abs_diff'] = max(out['m_max_abs_diff'], float(np.nanmax(np.abs(a - b))))
        out['rna_max_abs_diff'] = max(out['rna_max_abs_diff'], float(np.max(np.abs(D['rna'][di[s]] - R421['rna'][ri[s]]))))
        a, b = D['p'][di[s]], S['p'][si[s]]
        out['p_nan_equal'] &= bool(np.array_equal(np.isnan(a), np.isnan(b)))
        out['p_max_abs_diff'] = max(out['p_max_abs_diff'], float(np.nanmax(np.abs(a - b))) if np.isfinite(a).any() else 0.)
        out['tissue_equal'] &= bool(D['tissues'][di[s]] == S['tissues'][si[s]])
    out['passed'] = bool(ids and out['m_nan_equal'] and out['p_nan_equal'] and out['tissue_equal']
                         and max(out['m_max_abs_diff'], out['rna_max_abs_diff'], out['p_max_abs_diff']) <= 1e-9)
    return out


# ---------------------------------------------------------------- GDSC2 AUC (verbatim original_prepare_gdsc421.py::extract)
def gdsc_extract(source, ids):
    whitelist = set(ids); values = {}; conflicts = set(); names = defaultdict(set); targets = defaultdict(set)
    total = 0; matched = 0; invalid = 0; datasets = Counter(); releases = Counter(); duplicate_same = 0; outside_unit_auc = 0
    with open(source, encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        required = {'SANGER_MODEL_ID', 'DATASET', 'DRUG_ID', 'DRUG_NAME', 'WEBRELEASE', 'AUC', 'LN_IC50', 'MIN_CONC', 'MAX_CONC', 'PUTATIVE_TARGET'}
        if not required.issubset(reader.fieldnames):
            raise ValueError('unexpected source schema')
        for row in reader:
            total += 1; datasets[row['DATASET']] += 1; releases[row['WEBRELEASE']] += 1
            sid = row['SANGER_MODEL_ID']
            if sid not in whitelist:
                continue
            if row['DATASET'] != 'GDSC2':
                raise ValueError('mixed dataset')
            matched += 1; drug = row['DRUG_ID']; key = (sid, drug)
            names[drug].add(row['DRUG_NAME']); targets[drug].add(row['PUTATIVE_TARGET'])
            try:
                record = tuple(float(row[k]) for k in ('AUC', 'LN_IC50', 'MIN_CONC', 'MAX_CONC'))
            except ValueError:
                invalid += 1; continue
            if not np.isfinite(record).all() or not (0 < record[2] <= record[3]):
                invalid += 1; continue
            if not 0 <= record[0] <= 1:
                outside_unit_auc += 1
            if key in values:
                if values[key] != record:
                    conflicts.add(key)
                else:
                    duplicate_same += 1
            else:
                values[key] = record
    drugs = sorted(names, key=lambda x: int(x)); lookup = {v: i for i, v in enumerate(drugs)}; rows = {v: i for i, v in enumerate(ids)}
    auc = np.full((len(ids), len(drugs)), np.nan); ic50 = auc.copy()
    for (sid, drug), record in values.items():
        if (sid, drug) in conflicts:
            continue
        auc[rows[sid], lookup[drug]] = record[0]; ic50[rows[sid], lookup[drug]] = record[1]
    quality = dict(source_rows=total, whitelisted_source_rows=matched, datasets=dict(datasets), webreleases=dict(releases),
                   samples_requested=len(ids), samples_with_response=int(np.isfinite(auc).any(1).sum()), drug_ids=len(drugs),
                   observed_pairs=int(np.isfinite(auc).sum()), missing_pairs=int(np.isnan(auc).sum()), invalid_rows=invalid,
                   exact_duplicate_rows=duplicate_same, conflicting_keys=len(conflicts), outside_unit_auc_rows_retained=outside_unit_auc)
    return auc, ic50, np.array(drugs), quality


def gdsc_parity_421(source, S):
    """Re-derive the 421 AUC from `source` and compare with source-inputs (1e-9). Reads only 421 rows' values."""
    auc, _, drugs, q = gdsc_extract(source, S['sample_ids'].tolist())
    out = dict(source_sha256=sha(source), source_is_24jul22=None, drugs_equal=bool(np.array_equal(drugs, S['drug_ids'])),
               max_abs_diff=None, nan_pattern_equal=None, quality=q)
    out['source_is_24jul22'] = out['source_sha256'] == GDSC_24JUL22_SHA
    if out['drugs_equal']:
        out['max_abs_diff'] = float(np.nanmax(np.abs(auc - S['auc'])))
        out['nan_pattern_equal'] = bool(np.array_equal(np.isnan(auc), np.isnan(S['auc'])))
    out['passed'] = bool(out['drugs_equal'] and out['nan_pattern_equal'] and out['max_abs_diff'] <= 1e-9)
    return out


def gdsc_load_ids(source, ids, drug_ids):
    """AUC matrix (len(ids) x 295) on the frozen drug order for the given IDs; drugs outside the 295 are ignored and counted."""
    auc, _, drugs, q = gdsc_extract(source, list(ids))
    pos = {d: j for j, d in enumerate(drug_ids)}
    out = np.full((len(ids), len(drug_ids)), np.nan); extra = []
    for j, d in enumerate(drugs):
        if d in pos:
            out[:, pos[d]] = auc[:, j]
        else:
            extra.append(str(d))
    q['drugs_outside_frozen_295'] = extra; q['frozen_drugs_observed'] = int(np.isfinite(out).any(0).sum())
    return out, q


# ---------------------------------------------------------------- DepMap CRISPR gene effect (mirrors prepare_crispr_development.py)
def index_models(rows, safe_ids):
    """Verbatim creed/science/functional_source_catalog.py::index_models minus the alias table."""
    safe_ids = list(safe_ids)
    if len(safe_ids) != len(set(safe_ids)) or not safe_ids:
        raise ValueError('safe identity list must be nonempty and unique')
    if any(not re.fullmatch(r'SIDM\d+', s) for s in safe_ids):
        raise ValueError('only exact Sanger model IDs are accepted')
    rows = list(rows); safe = set(safe_ids)
    mids = Counter(r['ModelID'] for r in rows)
    sids = Counter(r['SangerModelID'] for r in rows if r['SangerModelID'])
    mapped = {}
    for row in rows:
        mid, sid = row['ModelID'], row['SangerModelID']
        if sid in safe and mids[mid] == 1 and sids[sid] == 1:
            mapped[mid] = {'sanger_id': sid, 'patient_id': row['PatientID'], 'lineage': row['OncotreeLineage'], 'cell_line_name': row['CellLineName']}
    return mapped


def crispr_extract(stream, mapping):
    """Verbatim prepare_crispr_development.py::extract — only rows whose ModelID is in `mapping` are decoded."""
    header = next(csv.reader([stream.readline().decode('utf-8-sig')]))
    if header[0] not in ('', 'ModelID') or len(set(header[1:])) != len(header) - 1:
        raise ValueError('invalid CRISPR header')
    records = {}; rows = 0
    for line in stream:
        prefix, delimiter, _ = line.partition(b',')
        if not delimiter:
            raise ValueError('invalid source identity delimiter')
        mid = prefix.decode('ascii').strip('"'); rows += 1
        if mid not in mapping:
            continue
        if mid in records:
            raise ValueError('duplicate selected outcome row')
        fields = next(csv.reader([line.decode('utf-8')]))
        if len(fields) != len(header):
            raise ValueError('invalid selected outcome width')
        values = np.array([np.nan if v in ('', 'NA', 'NaN', 'nan') else float(v) for v in fields[1:]])
        if np.isinf(values).any():
            raise ValueError('infinite selected outcome')
        records[mid] = values
    return header[1:], records, rows


def crispr_load(effect_path, model_path, ids, gene_ids, forbidden):
    """Gene-effect matrix (len(ids) x len(gene_ids)) for the given SIDMs via Model.csv exact SangerModelID; NaN rows for
    lines without a mapped/screened model. Returns (F, info)."""
    check_allowed(ids, forbidden)
    with open(model_path, encoding='utf-8-sig', newline='') as f:
        mapping = index_models(list(csv.DictReader(f)), list(ids))
    with open(effect_path, 'rb') as f:
        genes, records, nrows = crispr_extract(f, mapping)
    gpos = {g: j for j, g in enumerate(genes)}
    missing_genes = [g for g in gene_ids if g not in gpos]
    if missing_genes:
        raise ValueError(f'{len(missing_genes)} frozen CRISPR genes are missing from the effect file header')
    colsel = np.array([gpos[g] for g in gene_ids])
    F = np.full((len(ids), len(gene_ids)), np.nan); pos = {s: i for i, s in enumerate(ids)}; found = []
    for mid, vals in records.items():
        s = mapping[mid]['sanger_id']
        if pos[s] in found:
            raise ValueError('two screened models map to one allowed line')
        F[pos[s]] = vals[colsel]; found.append(pos[s])
    info = dict(effect_md5=md5(effect_path), effect_bytes=int(os.path.getsize(effect_path)),
                effect_is_24q2=None, model_sha256=sha(model_path), source_rows=nrows, header_genes=len(genes),
                mapped_models=len(mapping), lines_with_crispr=len(found), extra_header_genes=int(len(genes) - len(gene_ids)))
    info['effect_is_24q2'] = bool(info['effect_md5'] == DEPMAP_24Q2_EFFECT_MD5 and info['effect_bytes'] == DEPMAP_24Q2_EFFECT_BYTES)
    return F, info


def crispr_parity_310(effect_path, model_path, S, C):
    """Re-derive the 310 development lines' effects from the file and compare with crispr310-prepare effects.npy."""
    F, info = crispr_load(effect_path, model_path, S['sample_ids'].tolist(), C['gene_ids'].tolist(), forbidden=())
    pos = {s: i for i, s in enumerate(S['sample_ids'].tolist())}
    rows = np.array([pos[s] for s in C['sample_ids'].tolist()])
    sub = F[rows]
    out = dict(info=info, n_dev_lines_expected=int(len(rows)), n_dev_lines_found=int(np.isfinite(sub).any(1).sum()),
               lines_with_crispr_in_421=int(np.isfinite(F).any(1).sum()),
               nan_pattern_equal=bool(np.array_equal(np.isnan(sub), np.isnan(C['effects']))),
               max_abs_diff=float(np.nanmax(np.abs(sub - C['effects']))) if np.isfinite(sub).any() else None)
    out['passed'] = bool(out['nan_pattern_equal'] and out['max_abs_diff'] is not None and out['max_abs_diff'] <= 1e-9
                         and out['n_dev_lines_found'] == len(rows) and out['lines_with_crispr_in_421'] == len(rows))
    return out, F


def write_synthetic_gdsc(path, S):
    """DRYRUN ONLY: a GDSC2-format CSV holding the 421 development AUC values (no sealed value) to exercise the AUC reader."""
    with open(path, 'w', newline='') as fh:
        w = csv.writer(fh)
        w.writerow(['DATASET', 'SANGER_MODEL_ID', 'DRUG_ID', 'DRUG_NAME', 'WEBRELEASE', 'AUC', 'LN_IC50', 'MIN_CONC', 'MAX_CONC', 'PUTATIVE_TARGET'])
        for i, s in enumerate(S['sample_ids']):
            for j, d in enumerate(S['drug_ids']):
                v = S['auc'][i, j]
                if np.isfinite(v):
                    w.writerow(['GDSC2', s, d, f'drug{d}', 'Y', repr(float(v)), '0.0', '0.001', '10.0', 'synthetic'])


def write_synthetic_crispr(path, C):
    """DRYRUN ONLY: a CRISPRGeneEffect-format CSV with the 310 development lines (allowed values) to exercise the reader."""
    with open(path, 'w', newline='') as fh:
        w = csv.writer(fh)
        w.writerow([''] + C['gene_ids'].tolist())
        for i, mid in enumerate(C['model_ids'].tolist()):
            w.writerow([mid] + ['' if not np.isfinite(v) else repr(float(v)) for v in C['effects'][i]])

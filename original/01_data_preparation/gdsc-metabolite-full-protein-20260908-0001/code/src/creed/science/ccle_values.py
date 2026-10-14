"""Custody-safe materialisation and loading of frozen CCLE assay values.

This module is deliberately separate from :mod:`creed.science.ccle`, which is
metadata-only.  Values may be read only by the create-once custody CLI.  The
scientific loader below accepts the development-visible artifact and rejects
sealed artifacts and every sample absent from the visible partition.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import json
import math
import os
import re
from dataclasses import dataclass
from pathlib import Path
from types import MappingProxyType
from typing import Iterable, Mapping, Sequence

import numpy as np

from creed.science.ccle import (
    CCLEContractError,
    canonical_json_sha256,
    read_mmc2_identity_metadata,
    sha256_file,
    sha256_sorted_ids,
)


PROTOCOL_ID = "CREED_CCLE_L0_L1_L2_SPLIT_V1"
SPLIT_SCHEMA = "creed.ccle-l0-l1-l2-split.v1"
FROZEN_REAL_SPLIT_PAYLOAD_SHA256 = (
    "c29aa1ddc02c6f7a76c418e2f1dcb8124e3378031c49fc941bb6d729aba5dcf3"
)
VISIBLE_MANIFEST_SCHEMA = "creed.ccle-visible-materialization.v1"
SEALED_MANIFEST_SCHEMA = "creed.ccle-sealed-materialization.v1"
VISIBLE_ARTIFACT_ROLE = "development_visible"
SEALED_ARTIFACT_ROLE = "sealed_internal_final"
VISIBLE_PARTITIONS = ("train", "validation")
SOURCE_ROLES = (
    "L0_metabolomics",
    "L1_rna_expected_counts",
    "identity_mmc2",
    "L2_protein_mmc3",
)
EXPECTED_REAL_FEATURE_COUNTS = MappingProxyType(
    {
        "L0_metabolomics": 225,
        "L1_rna_expected_counts": 37_279,
        "L2_protein_mmc3": 6_692,
    }
)
_SIDM_PATTERN = re.compile(r"^SIDM[0-9]+$")
_MISSING_TOKENS = frozenset({"", "na", "nan", "null", "none", "#n/a"})


class CCLEValueCustodyError(ValueError):
    """Raised when materialisation or visible loading violates custody."""


def _ordered_text_sha256(values: Iterable[str]) -> str:
    digest = hashlib.sha256()
    for value in values:
        digest.update(value.encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest()


def _require_sidm(value: str, *, context: str) -> str:
    value = value.strip()
    if _SIDM_PATTERN.fullmatch(value) is None:
        raise CCLEValueCustodyError(f"invalid SIDM in {context}")
    return value


def _require_unique(values: Sequence[str], *, context: str) -> None:
    if len(values) != len(set(values)):
        raise CCLEValueCustodyError(f"duplicate identifiers in {context}")


def _parse_numeric(value: object, *, context: str, nonnegative: bool = False) -> float:
    if value is None:
        return math.nan
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.lower() in _MISSING_TOKENS:
            return math.nan
        candidate: object = stripped
    else:
        candidate = value
    try:
        parsed = float(candidate)
    except (TypeError, ValueError) as exc:
        raise CCLEValueCustodyError(f"non-numeric assay entry at {context}") from exc
    if math.isinf(parsed):
        raise CCLEValueCustodyError(f"infinite assay entry at {context}")
    if nonnegative and not math.isnan(parsed) and parsed < 0:
        raise CCLEValueCustodyError(f"negative expected count at {context}")
    return parsed


def _validate_partition(manifest: Mapping[str, object], name: str) -> tuple[str, ...]:
    partitions = manifest.get("partitions")
    if not isinstance(partitions, Mapping) or name not in partitions:
        raise CCLEValueCustodyError(f"split missing partition {name}")
    payload = partitions[name]
    if not isinstance(payload, Mapping):
        raise CCLEValueCustodyError(f"invalid split partition {name}")
    raw_ids = payload.get("sample_ids")
    if not isinstance(raw_ids, list) or not all(
        isinstance(item, str) for item in raw_ids
    ):
        raise CCLEValueCustodyError(f"invalid sample identifiers in partition {name}")
    sample_ids = tuple(_require_sidm(item, context=name) for item in raw_ids)
    _require_unique(sample_ids, context=f"partition {name}")
    if payload.get("count") != len(sample_ids):
        raise CCLEValueCustodyError(f"partition count mismatch for {name}")
    if payload.get("sample_ids_sha256") != sha256_sorted_ids(sample_ids):
        raise CCLEValueCustodyError(f"partition identifier hash mismatch for {name}")
    return sample_ids


@dataclass(frozen=True, slots=True)
class FrozenCCLESplit:
    split_path: Path
    split_file_sha256: str
    split_payload_sha256: str
    source_fingerprints: Mapping[str, Mapping[str, object]]
    train_sample_ids: tuple[str, ...]
    validation_sample_ids: tuple[str, ...]
    internal_final_sample_ids: tuple[str, ...]

    @property
    def visible_sample_ids(self) -> tuple[str, ...]:
        return self.train_sample_ids + self.validation_sample_ids

    @property
    def all_sample_ids(self) -> tuple[str, ...]:
        return self.visible_sample_ids + self.internal_final_sample_ids


def validate_frozen_split(
    split_path: str | Path, *, require_real_payload: bool = False
) -> FrozenCCLESplit:
    """Validate the split independently of all assay-value artifacts."""

    split_source = Path(split_path)
    if not split_source.is_file():
        raise CCLEValueCustodyError("frozen split file is missing")
    try:
        manifest = json.loads(split_source.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise CCLEValueCustodyError("frozen split is not valid UTF-8 JSON") from exc
    if not isinstance(manifest, dict):
        raise CCLEValueCustodyError("frozen split must be a JSON object")
    if (
        manifest.get("schema") != SPLIT_SCHEMA
        or manifest.get("protocol_id") != PROTOCOL_ID
    ):
        raise CCLEValueCustodyError("unexpected frozen split identity")
    integrity = manifest.get("integrity")
    if not isinstance(integrity, Mapping):
        raise CCLEValueCustodyError("frozen split lacks integrity record")
    expected_payload_hash = integrity.get("canonical_payload_sha256")
    payload = dict(manifest)
    payload.pop("integrity", None)
    observed_payload_hash = canonical_json_sha256(payload)
    if expected_payload_hash != observed_payload_hash:
        raise CCLEValueCustodyError("frozen split payload hash mismatch")
    if (
        require_real_payload
        and observed_payload_hash != FROZEN_REAL_SPLIT_PAYLOAD_SHA256
    ):
        raise CCLEValueCustodyError("scientific loader requires the frozen real split")

    release_gate = manifest.get("internal_final_release_gate")
    if (
        not isinstance(release_gate, Mapping)
        or release_gate.get("released") is not False
    ):
        raise CCLEValueCustodyError("internal-final release gate is not frozen closed")

    train = _validate_partition(manifest, "train")
    validation = _validate_partition(manifest, "validation")
    internal_final = _validate_partition(manifest, "internal_final")
    partition_sets = tuple(map(set, (train, validation, internal_final)))
    if any(
        partition_sets[left] & partition_sets[right]
        for left in range(3)
        for right in range(left + 1, 3)
    ):
        raise CCLEValueCustodyError("frozen split partitions overlap")
    universe = set().union(*partition_sets)
    universe_payload = manifest.get("sample_universe")
    if not isinstance(universe_payload, Mapping):
        raise CCLEValueCustodyError("frozen split lacks sample universe")
    if universe_payload.get("exact_triple_samples") != len(universe):
        raise CCLEValueCustodyError("frozen split universe count mismatch")
    if universe_payload.get("exact_triple_sample_ids_sha256") != sha256_sorted_ids(
        universe
    ):
        raise CCLEValueCustodyError("frozen split universe hash mismatch")

    source_manifest = manifest.get("sources")
    if not isinstance(source_manifest, Mapping) or set(source_manifest) != set(
        SOURCE_ROLES
    ):
        raise CCLEValueCustodyError("frozen split lacks source fingerprints")
    source_fingerprints: dict[str, Mapping[str, object]] = {}
    for role in SOURCE_ROLES:
        expected = source_manifest.get(role)
        if not isinstance(expected, Mapping):
            raise CCLEValueCustodyError(f"frozen split lacks source role {role}")
        source_fingerprints[role] = MappingProxyType(dict(expected))

    return FrozenCCLESplit(
        split_path=split_source.resolve(),
        split_file_sha256=sha256_file(split_source),
        split_payload_sha256=observed_payload_hash,
        source_fingerprints=MappingProxyType(source_fingerprints),
        train_sample_ids=train,
        validation_sample_ids=validation,
        internal_final_sample_ids=internal_final,
    )


@dataclass(frozen=True, slots=True)
class FrozenCCLEValueContract:
    split_path: Path
    split_file_sha256: str
    split_payload_sha256: str
    source_paths: Mapping[str, Path]
    source_sha256: Mapping[str, str]
    source_bytes: Mapping[str, int]
    train_sample_ids: tuple[str, ...]
    validation_sample_ids: tuple[str, ...]
    internal_final_sample_ids: tuple[str, ...]

    @property
    def visible_sample_ids(self) -> tuple[str, ...]:
        return self.train_sample_ids + self.validation_sample_ids

    @property
    def all_sample_ids(self) -> tuple[str, ...]:
        return self.visible_sample_ids + self.internal_final_sample_ids


def validate_frozen_split_and_sources(
    split_path: str | Path,
    *,
    source_paths: Mapping[str, str | Path],
) -> FrozenCCLEValueContract:
    """Verify the frozen split and every full source hash before value access."""

    split = validate_frozen_split(split_path)

    if set(source_paths) != set(SOURCE_ROLES):
        raise CCLEValueCustodyError("source roles do not match frozen contract")
    resolved_paths: dict[str, Path] = {}
    hashes: dict[str, str] = {}
    sizes: dict[str, int] = {}
    for role in SOURCE_ROLES:
        supplied = Path(source_paths[role])
        expected = split.source_fingerprints[role]
        if not supplied.is_file():
            raise CCLEValueCustodyError(f"source file missing for role {role}")
        observed_size = supplied.stat().st_size
        observed_hash = sha256_file(supplied)
        if (
            expected.get("bytes") != observed_size
            or expected.get("sha256") != observed_hash
        ):
            raise CCLEValueCustodyError(f"source fingerprint mismatch for role {role}")
        resolved_paths[role] = supplied.resolve()
        hashes[role] = observed_hash
        sizes[role] = observed_size

    rna_name = resolved_paths["L1_rna_expected_counts"].name.lower()
    if (
        not rna_name.endswith(".csv.gz")
        or "18q2" in rna_name
        or "rpkm" in rna_name
        or rna_name.endswith(".gct")
    ):
        raise CCLEValueCustodyError("RNA source is legacy, truncated, or derived")

    return FrozenCCLEValueContract(
        split_path=split.split_path,
        split_file_sha256=split.split_file_sha256,
        split_payload_sha256=split.split_payload_sha256,
        source_paths=MappingProxyType(resolved_paths),
        source_sha256=MappingProxyType(hashes),
        source_bytes=MappingProxyType(sizes),
        train_sample_ids=split.train_sample_ids,
        validation_sample_ids=split.validation_sample_ids,
        internal_final_sample_ids=split.internal_final_sample_ids,
    )


@dataclass(frozen=True, slots=True)
class CCLESourceValues:
    sample_ids: tuple[str, ...]
    l0_feature_names: tuple[str, ...]
    l0_values: np.ndarray
    l1_gene_ids: tuple[str, ...]
    l1_symbols: tuple[str, ...]
    l1_values: np.ndarray
    l2_protein_ids: tuple[str, ...]
    l2_values: np.ndarray


def _metabolomics_bridge(mmc2_path: Path) -> tuple[dict[str, str], dict[str, str]]:
    try:
        identities = read_mmc2_identity_metadata(mmc2_path)
    except CCLEContractError as exc:
        raise CCLEValueCustodyError("mmc2 identity bridge failed validation") from exc
    by_ccle: dict[str, str] = {}
    by_broad: dict[str, str] = {}
    for sidm, record in identities.items():
        if record.ccle_id is not None:
            if record.ccle_id in by_ccle:
                raise CCLEValueCustodyError("mmc2 CCLE_ID bridge is not one-to-one")
            by_ccle[record.ccle_id] = sidm
        if record.broad_id is not None:
            if record.broad_id in by_broad:
                raise CCLEValueCustodyError("mmc2 BROAD_ID bridge is not one-to-one")
            by_broad[record.broad_id] = sidm
    return by_ccle, by_broad


def read_l0_metabolomics_values(
    path: str | Path,
    mmc2_path: str | Path,
    sample_ids: Sequence[str],
    *,
    expected_feature_count: int,
) -> tuple[tuple[str, ...], np.ndarray]:
    """Read source L0 values in frozen SIDM order without filling missing data."""

    required = tuple(sample_ids)
    _require_unique(required, context="L0 requested samples")
    by_ccle, by_broad = _metabolomics_bridge(Path(mmc2_path))
    selected: dict[str, np.ndarray] = {}
    with Path(path).open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            header = next(reader)
        except StopIteration as exc:
            raise CCLEValueCustodyError("empty L0 source") from exc
        if header[:2] != ["CCLE_ID", "DepMap_ID"]:
            raise CCLEValueCustodyError("unexpected L0 source identity columns")
        features = tuple(item.strip() for item in header[2:])
        if len(features) != expected_feature_count or any(
            not item for item in features
        ):
            raise CCLEValueCustodyError("unexpected L0 feature count or identity")
        _require_unique(features, context="L0 feature names")
        required_set = set(required)
        for row_number, row in enumerate(reader, start=2):
            if not row or not any(item.strip() for item in row):
                continue
            if len(row) != len(header):
                raise CCLEValueCustodyError(f"L0 width mismatch at row {row_number}")
            ccle_id = row[0].strip()
            broad_id = row[1].strip()
            from_ccle = by_ccle.get(ccle_id) if ccle_id else None
            from_broad = by_broad.get(broad_id) if broad_id else None
            if (
                from_ccle is not None
                and from_broad is not None
                and from_ccle != from_broad
            ):
                raise CCLEValueCustodyError("conflicting L0 identity bridge")
            sidm = from_ccle or from_broad
            if sidm not in required_set:
                continue
            if sidm in selected:
                raise CCLEValueCustodyError("duplicate L0 row for frozen sample")
            selected[sidm] = np.fromiter(
                (
                    _parse_numeric(
                        value, context=f"L0 row {row_number} column {column}"
                    )
                    for column, value in enumerate(row[2:], start=3)
                ),
                dtype=np.float64,
                count=len(features),
            )
    missing = [sample_id for sample_id in required if sample_id not in selected]
    if missing:
        raise CCLEValueCustodyError("L0 source lacks one or more frozen samples")
    return features, np.vstack([selected[sample_id] for sample_id in required])


def read_l1_expected_count_values(
    path: str | Path,
    sample_ids: Sequence[str],
    *,
    expected_feature_count: int,
) -> tuple[tuple[str, ...], tuple[str, ...], np.ndarray]:
    """Read complete source expected counts using gene_id as the primary key."""

    required = tuple(sample_ids)
    _require_unique(required, context="L1 requested samples")
    source = Path(path)
    lower_name = source.name.lower()
    if (
        not lower_name.endswith(".csv.gz")
        or "18q2" in lower_name
        or "rpkm" in lower_name
    ):
        raise CCLEValueCustodyError("legacy/truncated RNA source is forbidden")
    with gzip.open(source, "rt", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        schema_rows: list[list[str]] = []
        for row_number in range(1, 5):
            try:
                schema_rows.append(next(reader))
            except StopIteration as exc:
                raise CCLEValueCustodyError(
                    f"RNA source ended before schema row {row_number}"
                ) from exc
        if [row[0] if row else None for row in schema_rows] != [
            "model_id",
            "model_name",
            "dataset_name",
            "gene_id",
        ]:
            raise CCLEValueCustodyError("unexpected RNA four-row schema")
        width = len(schema_rows[0])
        if width < 3 or any(len(row) != width for row in schema_rows):
            raise CCLEValueCustodyError("inconsistent RNA schema widths")
        if schema_rows[3][1].strip() != "symbol":
            raise CCLEValueCustodyError("RNA primary feature identity is not gene_id")
        source_sample_ids = tuple(
            _require_sidm(value, context="RNA model_id header")
            for value in schema_rows[0][2:]
        )
        _require_unique(source_sample_ids, context="RNA model_id header")
        index_by_sample = {
            value: index + 2 for index, value in enumerate(source_sample_ids)
        }
        if any(sample_id not in index_by_sample for sample_id in required):
            raise CCLEValueCustodyError("RNA source lacks one or more frozen samples")
        selected_columns = tuple(index_by_sample[sample_id] for sample_id in required)
        matrix = np.empty((len(required), expected_feature_count), dtype=np.float64)
        gene_ids: list[str] = []
        symbols: list[str] = []
        for source_row_number, row in enumerate(reader, start=5):
            if not row or not any(value.strip() for value in row):
                continue
            gene_index = len(gene_ids)
            if gene_index >= expected_feature_count:
                raise CCLEValueCustodyError(
                    "RNA source has more than expected features"
                )
            if len(row) != width:
                raise CCLEValueCustodyError(
                    f"RNA width mismatch at row {source_row_number}"
                )
            gene_id = row[0].strip()
            if not gene_id:
                raise CCLEValueCustodyError(
                    f"empty RNA gene_id at row {source_row_number}"
                )
            gene_ids.append(gene_id)
            symbols.append(row[1].strip())
            for sample_index, source_column in enumerate(selected_columns):
                matrix[sample_index, gene_index] = _parse_numeric(
                    row[source_column],
                    context=f"RNA row {source_row_number} column {source_column + 1}",
                    nonnegative=True,
                )
    if len(gene_ids) != expected_feature_count:
        raise CCLEValueCustodyError("unexpected RNA feature count")
    _require_unique(gene_ids, context="RNA gene_id primary keys")
    return tuple(gene_ids), tuple(symbols), matrix


def _openpyxl():
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - runtime-specific
        raise CCLEValueCustodyError("openpyxl is required for source mmc3") from exc
    return openpyxl


def read_l2_protein_values(
    path: str | Path,
    sample_ids: Sequence[str],
    *,
    expected_feature_count: int,
    sheet_name: str = "Prot matrix excl single-peptide",
) -> tuple[tuple[str, ...], np.ndarray]:
    """Read the frozen mmc3 protein sheet and preserve source protein IDs."""

    required = tuple(sample_ids)
    _require_unique(required, context="L2 requested samples")
    openpyxl = _openpyxl()
    workbook = openpyxl.load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise CCLEValueCustodyError("required mmc3 protein sheet is missing")
        sheet = workbook[sheet_name]
        try:
            header = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        except StopIteration as exc:
            raise CCLEValueCustodyError("mmc3 protein sheet lacks header") from exc
        if not header or str(header[0]).strip() != "Project_Identifier":
            raise CCLEValueCustodyError("unexpected mmc3 protein identity column")
        proteins = tuple(
            "" if value is None else str(value).strip() for value in header[1:]
        )
        while proteins and not proteins[-1]:
            proteins = proteins[:-1]
        if len(proteins) != expected_feature_count or any(
            not item for item in proteins
        ):
            raise CCLEValueCustodyError(
                "unexpected mmc3 protein feature count or identity"
            )
        _require_unique(proteins, context="mmc3 protein IDs")
        required_set = set(required)
        selected: dict[str, np.ndarray] = {}
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=3, max_col=len(proteins) + 1, values_only=True),
            start=3,
        ):
            if not row or row[0] is None:
                continue
            raw_id = str(row[0]).strip()
            sidm = _require_sidm(raw_id.split(";", 1)[0], context="mmc3 row")
            if sidm not in required_set:
                continue
            if sidm in selected:
                raise CCLEValueCustodyError("duplicate mmc3 row for frozen sample")
            selected[sidm] = np.fromiter(
                (
                    _parse_numeric(
                        value, context=f"L2 row {row_number} column {column}"
                    )
                    for column, value in enumerate(row[1:], start=2)
                ),
                dtype=np.float64,
                count=len(proteins),
            )
    finally:
        workbook.close()
    if any(sample_id not in selected for sample_id in required):
        raise CCLEValueCustodyError("mmc3 source lacks one or more frozen samples")
    return proteins, np.vstack([selected[sample_id] for sample_id in required])


def read_frozen_ccle_values(
    contract: FrozenCCLEValueContract,
    *,
    expected_feature_counts: Mapping[str, int] = EXPECTED_REAL_FEATURE_COUNTS,
    protein_sheet: str = "Prot matrix excl single-peptide",
) -> CCLESourceValues:
    """Read all frozen samples once, only after fingerprint verification."""

    if set(expected_feature_counts) != {
        "L0_metabolomics",
        "L1_rna_expected_counts",
        "L2_protein_mmc3",
    }:
        raise CCLEValueCustodyError("incomplete expected feature-count contract")
    sample_ids = contract.all_sample_ids
    l0_names, l0 = read_l0_metabolomics_values(
        contract.source_paths["L0_metabolomics"],
        contract.source_paths["identity_mmc2"],
        sample_ids,
        expected_feature_count=expected_feature_counts["L0_metabolomics"],
    )
    l1_gene_ids, l1_symbols, l1 = read_l1_expected_count_values(
        contract.source_paths["L1_rna_expected_counts"],
        sample_ids,
        expected_feature_count=expected_feature_counts["L1_rna_expected_counts"],
    )
    l2_ids, l2 = read_l2_protein_values(
        contract.source_paths["L2_protein_mmc3"],
        sample_ids,
        expected_feature_count=expected_feature_counts["L2_protein_mmc3"],
        sheet_name=protein_sheet,
    )
    return CCLESourceValues(
        sample_ids=sample_ids,
        l0_feature_names=l0_names,
        l0_values=l0,
        l1_gene_ids=l1_gene_ids,
        l1_symbols=l1_symbols,
        l1_values=l1,
        l2_protein_ids=l2_ids,
        l2_values=l2,
    )


def partition_arrays(
    values: CCLESourceValues,
    sample_ids: Sequence[str],
    partition_labels: Sequence[str],
) -> dict[str, np.ndarray]:
    """Select an exact ordered partition without transforming or imputing values."""

    requested = tuple(sample_ids)
    labels = tuple(partition_labels)
    if len(requested) != len(labels):
        raise CCLEValueCustodyError("sample and partition-label counts differ")
    _require_unique(requested, context="materialized partition")
    index = {
        sample_id: position for position, sample_id in enumerate(values.sample_ids)
    }
    if any(sample_id not in index for sample_id in requested):
        raise CCLEValueCustodyError("materialized partition includes unknown sample")
    rows = np.asarray([index[sample_id] for sample_id in requested], dtype=np.int64)
    return {
        "sample_ids": np.asarray(requested, dtype=np.str_),
        "partition_labels": np.asarray(labels, dtype=np.str_),
        "l0_feature_names": np.asarray(values.l0_feature_names, dtype=np.str_),
        "l0_values": values.l0_values[rows].copy(),
        "l1_gene_ids": np.asarray(values.l1_gene_ids, dtype=np.str_),
        "l1_symbols": np.asarray(values.l1_symbols, dtype=np.str_),
        "l1_values": values.l1_values[rows].copy(),
        "l2_protein_ids": np.asarray(values.l2_protein_ids, dtype=np.str_),
        "l2_values": values.l2_values[rows].copy(),
    }


def write_npz_create_once(path: str | Path, arrays: Mapping[str, np.ndarray]) -> None:
    """Write an NPZ through an O_EXCL file descriptor and fsync it."""

    destination = Path(path)
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
    descriptor = os.open(destination, flags, 0o600)
    try:
        with os.fdopen(descriptor, "wb", closefd=False) as handle:
            np.savez_compressed(handle, **arrays)
            handle.flush()
            os.fsync(handle.fileno())
    finally:
        os.close(descriptor)


def array_manifest(arrays: Mapping[str, np.ndarray]) -> dict[str, object]:
    """Return non-scientific structural hashes for the materialized arrays."""

    sample_ids = tuple(str(value) for value in arrays["sample_ids"].tolist())
    labels = tuple(str(value) for value in arrays["partition_labels"].tolist())
    return {
        "sample_count": len(sample_ids),
        "sample_ids_sha256": sha256_sorted_ids(sample_ids),
        "ordered_sample_ids_sha256": _ordered_text_sha256(sample_ids),
        "partition_labels_sha256": _ordered_text_sha256(labels),
        "layers": {
            "L0_metabolomics": {
                "shape": list(arrays["l0_values"].shape),
                "dtype": str(arrays["l0_values"].dtype),
                "feature_count": int(arrays["l0_feature_names"].shape[0]),
                "ordered_feature_ids_sha256": _ordered_text_sha256(
                    str(value) for value in arrays["l0_feature_names"].tolist()
                ),
                "missing_values_preserved": True,
            },
            "L1_rna_expected_counts": {
                "shape": list(arrays["l1_values"].shape),
                "dtype": str(arrays["l1_values"].dtype),
                "feature_count": int(arrays["l1_gene_ids"].shape[0]),
                "primary_key": "source_gene_id",
                "ordered_gene_ids_sha256": _ordered_text_sha256(
                    str(value) for value in arrays["l1_gene_ids"].tolist()
                ),
                "ordered_symbol_annotations_sha256": _ordered_text_sha256(
                    str(value) for value in arrays["l1_symbols"].tolist()
                ),
                "semantic": "untransformed_expected_counts",
                "missing_values_preserved": True,
            },
            "L2_protein_mmc3": {
                "shape": list(arrays["l2_values"].shape),
                "dtype": str(arrays["l2_values"].dtype),
                "feature_count": int(arrays["l2_protein_ids"].shape[0]),
                "ordered_protein_ids_sha256": _ordered_text_sha256(
                    str(value) for value in arrays["l2_protein_ids"].tolist()
                ),
                "source_sheet": "Prot matrix excl single-peptide",
                "missing_values_preserved": True,
            },
        },
        "missing_value_policy": "preserve_nan_never_fill_zero",
    }


@dataclass(frozen=True, slots=True)
class CCLEVisibleData:
    sample_ids: tuple[str, ...]
    partition_labels: tuple[str, ...]
    l0_feature_names: tuple[str, ...]
    l0_values: np.ndarray
    l1_gene_ids: tuple[str, ...]
    l1_symbols: tuple[str, ...]
    l1_values: np.ndarray
    l2_protein_ids: tuple[str, ...]
    l2_values: np.ndarray


_EXPECTED_NPZ_KEYS = frozenset(
    {
        "sample_ids",
        "partition_labels",
        "l0_feature_names",
        "l0_values",
        "l1_gene_ids",
        "l1_symbols",
        "l1_values",
        "l2_protein_ids",
        "l2_values",
    }
)


def _reject_sealed_path(path: Path) -> None:
    if any(
        "sealed" in part.lower() or "internal_final" in part.lower()
        for part in path.parts
    ):
        raise CCLEValueCustodyError(
            "scientific loader refuses sealed/internal-final paths"
        )


def _immutable(array: np.ndarray) -> np.ndarray:
    result = np.asarray(array)
    result.setflags(write=False)
    return result


def _validate_loaded_arrays(arrays: Mapping[str, np.ndarray]) -> None:
    vector_keys = (
        "sample_ids",
        "partition_labels",
        "l0_feature_names",
        "l1_gene_ids",
        "l1_symbols",
        "l2_protein_ids",
    )
    matrix_keys = ("l0_values", "l1_values", "l2_values")
    if any(arrays[key].ndim != 1 for key in vector_keys):
        raise CCLEValueCustodyError("visible identifier arrays must be one-dimensional")
    if any(arrays[key].ndim != 2 for key in matrix_keys):
        raise CCLEValueCustodyError("visible value arrays must be two-dimensional")
    sample_count = arrays["sample_ids"].shape[0]
    if arrays["partition_labels"].shape[0] != sample_count:
        raise CCLEValueCustodyError("visible partition-label shape mismatch")
    for layer, feature_key in (
        ("l0_values", "l0_feature_names"),
        ("l1_values", "l1_gene_ids"),
        ("l2_values", "l2_protein_ids"),
    ):
        matrix = arrays[layer]
        if matrix.shape != (sample_count, arrays[feature_key].shape[0]):
            raise CCLEValueCustodyError(f"visible {layer} shape mismatch")
        if matrix.dtype != np.dtype(np.float64):
            raise CCLEValueCustodyError(f"visible {layer} dtype mismatch")
        if np.isinf(matrix).any():
            raise CCLEValueCustodyError(f"visible {layer} contains infinity")
    if arrays["l1_symbols"].shape != arrays["l1_gene_ids"].shape:
        raise CCLEValueCustodyError("visible RNA symbol annotation shape mismatch")
    finite_rna = arrays["l1_values"][~np.isnan(arrays["l1_values"])]
    if np.any(finite_rna < 0):
        raise CCLEValueCustodyError("visible RNA expected counts contain negatives")
    sample_ids = tuple(str(value) for value in arrays["sample_ids"].tolist())
    _require_unique(sample_ids, context="visible artifact samples")
    for sample_id in sample_ids:
        _require_sidm(sample_id, context="visible artifact")
    for key in ("l0_feature_names", "l1_gene_ids", "l2_protein_ids"):
        identifiers = tuple(str(value) for value in arrays[key].tolist())
        if any(not value for value in identifiers):
            raise CCLEValueCustodyError(f"visible {key} contains empty identity")
        _require_unique(identifiers, context=f"visible {key}")


def load_visible_ccle_artifact(
    artifact_path: str | Path,
    *,
    split_path: str | Path,
    manifest_path: str | Path | None = None,
    requested_sample_ids: Sequence[str] | None = None,
) -> CCLEVisibleData:
    """Load only the hash-bound train/validation artifact, failing closed."""

    split = validate_frozen_split(split_path, require_real_payload=True)
    artifact = Path(artifact_path).resolve()
    manifest_source = (
        Path(manifest_path).resolve()
        if manifest_path is not None
        else artifact.with_name("manifest.json")
    )
    _reject_sealed_path(artifact)
    _reject_sealed_path(manifest_source)
    if not artifact.is_file() or not manifest_source.is_file():
        raise CCLEValueCustodyError("visible artifact or manifest is missing")
    try:
        manifest = json.loads(manifest_source.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise CCLEValueCustodyError("visible manifest is invalid") from exc
    if not isinstance(manifest, dict):
        raise CCLEValueCustodyError("visible manifest must be an object")
    integrity = manifest.get("integrity")
    if not isinstance(integrity, Mapping):
        raise CCLEValueCustodyError("visible manifest lacks integrity")
    payload = dict(manifest)
    payload.pop("integrity", None)
    if integrity.get("canonical_payload_sha256") != canonical_json_sha256(payload):
        raise CCLEValueCustodyError("visible manifest payload hash mismatch")
    if (
        manifest.get("schema") != VISIBLE_MANIFEST_SCHEMA
        or manifest.get("artifact_role") != VISIBLE_ARTIFACT_ROLE
        or manifest.get("internal_final_released") is not False
    ):
        raise CCLEValueCustodyError("manifest is not development-visible")
    expected_split_record = {
        "file_sha256": split.split_file_sha256,
        "payload_sha256": split.split_payload_sha256,
        "visible_order": "frozen train sample_ids then frozen validation sample_ids",
        "partition_counts": {
            "train": len(split.train_sample_ids),
            "validation": len(split.validation_sample_ids),
        },
        "internal_final_count": len(split.internal_final_sample_ids),
        "internal_final_sample_ids_sha256": sha256_sorted_ids(
            split.internal_final_sample_ids
        ),
    }
    if manifest.get("split") != expected_split_record:
        raise CCLEValueCustodyError("visible manifest does not bind the frozen split")
    artifact_record = manifest.get("artifact")
    if not isinstance(artifact_record, Mapping):
        raise CCLEValueCustodyError("visible manifest lacks artifact record")
    if artifact_record.get("filename") != artifact.name:
        raise CCLEValueCustodyError("visible artifact filename mismatch")
    if artifact_record.get("bytes") != artifact.stat().st_size:
        raise CCLEValueCustodyError("visible artifact size mismatch")
    if artifact_record.get("sha256") != sha256_file(artifact):
        raise CCLEValueCustodyError("visible artifact hash mismatch")

    try:
        with np.load(artifact, allow_pickle=False) as archive:
            if set(archive.files) != set(_EXPECTED_NPZ_KEYS):
                raise CCLEValueCustodyError("visible artifact key set mismatch")
            arrays = {key: np.asarray(archive[key]).copy() for key in archive.files}
    except (OSError, ValueError) as exc:
        if isinstance(exc, CCLEValueCustodyError):
            raise
        raise CCLEValueCustodyError("visible artifact is not a safe NPZ") from exc

    _validate_loaded_arrays(arrays)
    structural = array_manifest(arrays)
    expected_structure = manifest.get("structure")
    if structural != expected_structure:
        raise CCLEValueCustodyError("visible artifact structure mismatch")
    sample_ids = tuple(str(value) for value in arrays["sample_ids"].tolist())
    labels = tuple(str(value) for value in arrays["partition_labels"].tolist())
    if any(label not in VISIBLE_PARTITIONS for label in labels):
        raise CCLEValueCustodyError("visible artifact contains a forbidden partition")
    expected_ids = split.visible_sample_ids
    expected_labels = ("train",) * len(split.train_sample_ids) + ("validation",) * len(
        split.validation_sample_ids
    )
    if sample_ids != expected_ids or labels != expected_labels:
        raise CCLEValueCustodyError(
            "visible artifact differs from frozen train/validation order"
        )

    if requested_sample_ids is None:
        selected_ids = sample_ids
        row_indices = np.arange(len(sample_ids), dtype=np.int64)
    else:
        selected_ids = tuple(
            _require_sidm(value, context="scientific loader request")
            for value in requested_sample_ids
        )
        _require_unique(selected_ids, context="scientific loader request")
        index = {sample_id: position for position, sample_id in enumerate(sample_ids)}
        if any(sample_id not in index for sample_id in selected_ids):
            raise CCLEValueCustodyError(
                "scientific loader request includes non-visible or internal-final sample"
            )
        row_indices = np.asarray([index[sample_id] for sample_id in selected_ids])

    return CCLEVisibleData(
        sample_ids=selected_ids,
        partition_labels=tuple(labels[index] for index in row_indices.tolist()),
        l0_feature_names=tuple(
            str(value) for value in arrays["l0_feature_names"].tolist()
        ),
        l0_values=_immutable(arrays["l0_values"][row_indices].copy()),
        l1_gene_ids=tuple(str(value) for value in arrays["l1_gene_ids"].tolist()),
        l1_symbols=tuple(str(value) for value in arrays["l1_symbols"].tolist()),
        l1_values=_immutable(arrays["l1_values"][row_indices].copy()),
        l2_protein_ids=tuple(str(value) for value in arrays["l2_protein_ids"].tolist()),
        l2_values=_immutable(arrays["l2_values"][row_indices].copy()),
    )

"""Metadata-only contract and split construction for CCLE L0-L1-L2.

This module intentionally does not load a metabolite, RNA, or protein value.
It is limited to complete-file hashing, source schemas, sample identifiers, and
the non-expression identity/lineage metadata in ProCan Table S1 (``mmc2``).
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Mapping


class CCLEContractError(ValueError):
    """Raised when a source or split violates the frozen CCLE contract."""


SIDM_PATTERN = re.compile(r"^SIDM[0-9]+$")
PARTITION_ORDER = ("train", "validation", "internal_final")
DEFAULT_PARTITION_RATIOS = {
    "train": 0.70,
    "validation": 0.10,
    "internal_final": 0.20,
}


def _clean(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_sorted_ids(values: Iterable[str]) -> str:
    """Hash a set as sorted UTF-8 identifiers, one identifier per LF line."""

    digest = hashlib.sha256()
    for value in sorted(set(values)):
        digest.update(value.encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest()


def canonical_json_sha256(value: object) -> str:
    payload = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _require_unique_nonempty(
    pairs: Iterable[tuple[str, str | None]], *, field: str
) -> dict[str, str]:
    owner_by_value: dict[str, str] = {}
    for owner, value in pairs:
        if value is None:
            continue
        prior = owner_by_value.get(value)
        if prior is not None and prior != owner:
            raise CCLEContractError(
                f"{field} is not one-to-one: {value!r} maps to {prior!r} and {owner!r}"
            )
        owner_by_value[value] = owner
    return owner_by_value


def _require_sidm(value: str, *, source: str) -> str:
    if SIDM_PATTERN.fullmatch(value) is None:
        raise CCLEContractError(f"invalid SIDM identifier in {source}: {value!r}")
    return value


def _openpyxl():
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - exercised by minimal runtimes
        raise CCLEContractError(
            "openpyxl is required only for metadata-only mmc2/mmc3 inspection"
        ) from exc
    return openpyxl


@dataclass(frozen=True, slots=True)
class CCLEIdentityRecord:
    model_id: str
    project_identifier: str | None
    cell_line: str | None
    source: str | None
    identifier: str | None
    tissue_type: str | None
    cancer_type: str | None
    cancer_subtype: str | None
    haem_lineage: str | None
    broad_id: str | None
    ccle_id: str | None


@dataclass(frozen=True, slots=True)
class SourceFingerprint:
    role: str
    path: str
    sha256: str
    bytes: int
    inspection_scope: str


@dataclass(frozen=True, slots=True)
class CCLETripleInventory:
    identities: Mapping[str, CCLEIdentityRecord]
    protein_sidms: frozenset[str]
    rna_sidms: frozenset[str]
    metabolomics_sidms: frozenset[str]
    triple_sidms: frozenset[str]
    source_fingerprints: Mapping[str, SourceFingerprint]
    metabolomics_rows: int
    metabolomics_unmapped_rows: int
    rna_header_samples: int
    protein_header_samples: int
    protein_sheet: str


def read_mmc2_identity_metadata(
    path: str | Path,
    *,
    sheet_name: str = "Cell line level sample info",
) -> dict[str, CCLEIdentityRecord]:
    """Read Table S1 identity and lineage metadata, never assay values."""

    source = Path(path)
    if not source.is_file():
        raise CCLEContractError(f"missing mmc2 source: {source}")
    openpyxl = _openpyxl()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise CCLEContractError(f"missing mmc2 sheet: {sheet_name!r}")
        sheet = workbook[sheet_name]
        headers = [
            _clean(cell.value)
            for cell in next(sheet.iter_rows(min_row=2, max_row=2))
        ]
        header_index = {
            name: index for index, name in enumerate(headers) if name is not None
        }
        required = (
            "model_id",
            "Project_Identifier",
            "Cell_line",
            "Source",
            "Identifier",
            "Tissue_type",
            "Cancer_type",
            "Cancer_subtype",
            "Haem_lineage",
            "BROAD_ID",
            "CCLE_ID",
        )
        missing = [name for name in required if name not in header_index]
        if missing:
            raise CCLEContractError(f"mmc2 missing required metadata fields: {missing}")
        max_column = max(header_index[name] for name in required) + 1
        records: dict[str, CCLEIdentityRecord] = {}
        for row in sheet.iter_rows(
            min_row=3, max_col=max_column, values_only=True
        ):
            values = {name: _clean(row[header_index[name]]) for name in required}
            model_id = values["model_id"]
            if model_id is None:
                continue
            _require_sidm(model_id, source="mmc2 model_id")
            if model_id in records:
                raise CCLEContractError(f"duplicate mmc2 model_id: {model_id}")
            records[model_id] = CCLEIdentityRecord(
                model_id=model_id,
                project_identifier=values["Project_Identifier"],
                cell_line=values["Cell_line"],
                source=values["Source"],
                identifier=values["Identifier"],
                tissue_type=values["Tissue_type"],
                cancer_type=values["Cancer_type"],
                cancer_subtype=values["Cancer_subtype"],
                haem_lineage=values["Haem_lineage"],
                broad_id=values["BROAD_ID"],
                ccle_id=values["CCLE_ID"],
            )
    finally:
        workbook.close()
    if not records:
        raise CCLEContractError("mmc2 contains no identity records")
    _require_unique_nonempty(
        ((record.model_id, record.broad_id) for record in records.values()),
        field="mmc2 BROAD_ID",
    )
    _require_unique_nonempty(
        ((record.model_id, record.ccle_id) for record in records.values()),
        field="mmc2 CCLE_ID",
    )
    return records


def read_mmc3_sample_ids(
    path: str | Path,
    *,
    sheet_name: str = "Prot matrix excl single-peptide",
) -> frozenset[str]:
    """Read only the mmc3 header cell and first-column sample identifiers."""

    source = Path(path)
    if not source.is_file():
        raise CCLEContractError(f"missing mmc3 source: {source}")
    openpyxl = _openpyxl()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise CCLEContractError(f"missing mmc3 sheet: {sheet_name!r}")
        sheet = workbook[sheet_name]
        header = _clean(
            next(
                sheet.iter_rows(
                    min_row=2,
                    max_row=2,
                    min_col=1,
                    max_col=1,
                    values_only=True,
                )
            )[0]
        )
        if header != "Project_Identifier":
            raise CCLEContractError(
                f"unexpected mmc3 sample identifier field: {header!r}"
            )
        sample_ids: list[str] = []
        for (raw_value,) in sheet.iter_rows(
            min_row=3, min_col=1, max_col=1, values_only=True
        ):
            raw = _clean(raw_value)
            if raw is None:
                continue
            sidm = _clean(raw.split(";", 1)[0])
            if sidm is None:
                raise CCLEContractError(f"empty SIDM prefix in mmc3 value: {raw!r}")
            sample_ids.append(_require_sidm(sidm, source="mmc3 Project_Identifier"))
    finally:
        workbook.close()
    if len(sample_ids) != len(set(sample_ids)):
        duplicates = sorted(
            value for value, count in Counter(sample_ids).items() if count > 1
        )
        raise CCLEContractError(f"duplicate SIDMs in mmc3: {duplicates[:5]}")
    if not sample_ids:
        raise CCLEContractError("mmc3 contains no sample identifiers")
    return frozenset(sample_ids)


def read_rna_header_sample_ids(path: str | Path) -> frozenset[str]:
    """Read exactly four schema/header rows from the expected-count source."""

    source = Path(path)
    lower_name = source.name.lower()
    if lower_name.endswith(".gct") or "18q2" in lower_name or "rpkm" in lower_name:
        raise CCLEContractError("legacy/truncated CCLE GCT is forbidden")
    if not lower_name.endswith(".csv.gz"):
        raise CCLEContractError(
            "RNA input must be the complete expected-count .csv.gz source"
        )
    if not source.is_file():
        raise CCLEContractError(f"missing RNA source: {source}")
    with gzip.open(source, mode="rt", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        rows: list[list[str]] = []
        for index in range(4):
            try:
                rows.append(next(reader))
            except StopIteration as exc:
                raise CCLEContractError(
                    f"RNA source ended before schema row {index + 1}"
                ) from exc
    if [row[0] if row else None for row in rows] != [
        "model_id",
        "model_name",
        "dataset_name",
        "gene_id",
    ]:
        raise CCLEContractError("unexpected RNA four-row schema")
    if len(rows[3]) < 2 or rows[3][1].strip() != "symbol":
        raise CCLEContractError(
            "RNA feature identity must be source gene_id; "
            "symbol-only derivatives are forbidden"
        )
    if len(rows[0]) < 3:
        raise CCLEContractError("RNA header contains no sample columns")
    sample_ids = [value.strip() for value in rows[0][2:] if value.strip()]
    for sample_id in sample_ids:
        _require_sidm(sample_id, source="RNA model_id header")
    if len(sample_ids) != len(set(sample_ids)):
        raise CCLEContractError("duplicate SIDMs in RNA model_id header")
    return frozenset(sample_ids)


def read_metabolomics_identity_rows(
    path: str | Path,
) -> tuple[tuple[str | None, str | None], ...]:
    """Retain only CCLE_ID and DepMap_ID from each metabolomics sample row."""

    source = Path(path)
    if source.suffix.lower() != ".csv":
        raise CCLEContractError("metabolomics source must be the source CSV")
    if not source.is_file():
        raise CCLEContractError(f"missing metabolomics source: {source}")
    rows: list[tuple[str | None, str | None]] = []
    with source.open(mode="r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            header = next(reader)
        except StopIteration as exc:
            raise CCLEContractError("empty metabolomics CSV") from exc
        if header[:2] != ["CCLE_ID", "DepMap_ID"]:
            raise CCLEContractError(
                "metabolomics identity fields must be first and exact: "
                "CCLE_ID, DepMap_ID"
            )
        for row_number, row in enumerate(reader, start=2):
            if not row or not any(value.strip() for value in row):
                continue
            if len(row) < 2:
                raise CCLEContractError(
                    f"metabolomics row {row_number} lacks both identity fields"
                )
            pair = (_clean(row[0]), _clean(row[1]))
            if pair == (None, None):
                raise CCLEContractError(
                    f"metabolomics row {row_number} has no sample identity"
                )
            rows.append(pair)
    if not rows:
        raise CCLEContractError("metabolomics source contains no sample rows")
    for index, label in ((0, "CCLE_ID"), (1, "DepMap_ID")):
        values = [row[index] for row in rows if row[index] is not None]
        if len(values) != len(set(values)):
            raise CCLEContractError(f"duplicate metabolomics {label}")
    return tuple(rows)


def _fingerprint(path: Path, role: str, scope: str) -> SourceFingerprint:
    return SourceFingerprint(
        role=role,
        path=str(path),
        sha256=sha256_file(path),
        bytes=path.stat().st_size,
        inspection_scope=scope,
    )


def load_ccle_triple_inventory(
    *,
    metabolomics_path: str | Path,
    rna_path: str | Path,
    mmc2_path: str | Path,
    mmc3_path: str | Path,
    protein_sheet: str = "Prot matrix excl single-peptide",
    expected_triple_count: int | None = None,
    minimum_triple_count: int = 1,
) -> CCLETripleInventory:
    """Build the exact L0-L1-L2 SIDM intersection without assay-value access."""

    paths = {
        "L0_metabolomics": Path(metabolomics_path),
        "L1_rna_expected_counts": Path(rna_path),
        "identity_mmc2": Path(mmc2_path),
        "L2_protein_mmc3": Path(mmc3_path),
    }
    identities = read_mmc2_identity_metadata(paths["identity_mmc2"])
    protein_sidms = read_mmc3_sample_ids(
        paths["L2_protein_mmc3"], sheet_name=protein_sheet
    )
    metadata_sidms = frozenset(identities)
    if protein_sidms != metadata_sidms:
        missing_protein = sorted(metadata_sidms - protein_sidms)
        missing_metadata = sorted(protein_sidms - metadata_sidms)
        raise CCLEContractError(
            "mmc2/mmc3 sample identity mismatch; "
            f"missing_protein={missing_protein[:5]}, "
            f"missing_metadata={missing_metadata[:5]}"
        )

    rna_sidms = read_rna_header_sample_ids(paths["L1_rna_expected_counts"])
    metabolomics_rows = read_metabolomics_identity_rows(paths["L0_metabolomics"])
    broad_to_sidm = _require_unique_nonempty(
        ((record.model_id, record.broad_id) for record in identities.values()),
        field="mmc2 BROAD_ID",
    )
    ccle_to_sidm = _require_unique_nonempty(
        ((record.model_id, record.ccle_id) for record in identities.values()),
        field="mmc2 CCLE_ID",
    )
    mapped_metabolomics: list[str] = []
    unmapped_rows = 0
    for ccle_id, depmap_id in metabolomics_rows:
        from_ccle = ccle_to_sidm.get(ccle_id) if ccle_id is not None else None
        from_broad = broad_to_sidm.get(depmap_id) if depmap_id is not None else None
        if from_ccle is not None and from_broad is not None and from_ccle != from_broad:
            raise CCLEContractError(
                "conflicting mmc2 identity bridge for metabolomics row: "
                f"CCLE_ID={ccle_id!r}->{from_ccle}, "
                f"DepMap_ID={depmap_id!r}->{from_broad}"
            )
        sidm = from_ccle or from_broad
        if sidm is None:
            unmapped_rows += 1
        else:
            mapped_metabolomics.append(sidm)
    if len(mapped_metabolomics) != len(set(mapped_metabolomics)):
        raise CCLEContractError(
            "multiple metabolomics rows resolve to the same mmc2 SIDM"
        )
    metabolomics_sidms = frozenset(mapped_metabolomics)
    triple_sidms = protein_sidms & rna_sidms & metabolomics_sidms
    if len(triple_sidms) < minimum_triple_count:
        raise CCLEContractError(
            f"only {len(triple_sidms)} exact triple-layer samples; "
            f"minimum is {minimum_triple_count}"
        )
    if expected_triple_count is not None and len(triple_sidms) != expected_triple_count:
        raise CCLEContractError(
            f"exact triple count changed: expected {expected_triple_count}, "
            f"observed {len(triple_sidms)}"
        )

    scopes = {
        "L0_metabolomics": "full-byte hash; header plus CCLE_ID/DepMap_ID only",
        "L1_rna_expected_counts": "full-byte hash; first four schema/header rows only",
        "identity_mmc2": "full-byte hash; Table S1 non-expression metadata only",
        "L2_protein_mmc3": (
            "full-byte hash; selected-sheet header and first-column sample IDs only"
        ),
    }
    fingerprints = {
        role: _fingerprint(path, role, scopes[role]) for role, path in paths.items()
    }
    return CCLETripleInventory(
        identities=identities,
        protein_sidms=protein_sidms,
        rna_sidms=rna_sidms,
        metabolomics_sidms=metabolomics_sidms,
        triple_sidms=triple_sidms,
        source_fingerprints=fingerprints,
        metabolomics_rows=len(metabolomics_rows),
        metabolomics_unmapped_rows=unmapped_rows,
        rna_header_samples=len(rna_sidms),
        protein_header_samples=len(protein_sidms),
        protein_sheet=protein_sheet,
    )


def _normalise_label(value: str | None) -> str | None:
    if value is None:
        return None
    normalised = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
    return normalised or None


_DERIVATIVE_SUFFIX = re.compile(
    r"(?:\s+(?:parent(?:al)?|resistant|res|derived|derivative|subclone|clone)"
    r"(?:\s*[0-9a-z]+)?)$"
)


def _conservative_family(value: str | None) -> tuple[str | None, bool]:
    normalised = _normalise_label(value)
    if normalised is None:
        return None, False
    stripped = _DERIVATIVE_SUFFIX.sub("", normalised).strip()
    return stripped, stripped != normalised and bool(stripped)


class _UnionFind:
    def __init__(self, values: Iterable[str]) -> None:
        self.parent = {value: value for value in values}

    def find(self, value: str) -> str:
        parent = self.parent[value]
        if parent != value:
            self.parent[value] = self.find(parent)
        return self.parent[value]

    def union(self, left: str, right: str) -> None:
        left_root = self.find(left)
        right_root = self.find(right)
        if left_root == right_root:
            return
        low, high = sorted((left_root, right_root))
        self.parent[high] = low


def infer_identity_groups(
    identities: Mapping[str, CCLEIdentityRecord], sample_ids: Iterable[str]
) -> tuple[dict[str, str], list[dict[str, object]]]:
    """Conservatively keep exact aliases and explicit derivatives together."""

    selected = sorted(set(sample_ids))
    missing = [sample_id for sample_id in selected if sample_id not in identities]
    if missing:
        raise CCLEContractError(f"split samples missing from mmc2: {missing[:5]}")
    union = _UnionFind(selected)
    evidence_by_pair: defaultdict[tuple[str, str], set[str]] = defaultdict(set)

    def connect_by_key(label: str, keyed: Iterable[tuple[str, str | None]]) -> None:
        buckets: defaultdict[str, list[str]] = defaultdict(list)
        for sample_id, key in keyed:
            if key is not None:
                buckets[key].append(sample_id)
        for key, members in buckets.items():
            if len(members) < 2:
                continue
            anchor = min(members)
            for member in sorted(members):
                if member != anchor:
                    union.union(anchor, member)
                    evidence_by_pair[(anchor, member)].add(f"{label}:{key}")

    connect_by_key(
        "shared_source_identifier",
        (
            (sample_id, _normalise_label(identities[sample_id].identifier))
            for sample_id in selected
        ),
    )
    connect_by_key(
        "exact_cell_line_name",
        (
            (sample_id, _normalise_label(identities[sample_id].cell_line))
            for sample_id in selected
        ),
    )
    family_rows = [
        (sample_id, *_conservative_family(identities[sample_id].cell_line))
        for sample_id in selected
    ]
    derivative_families = {
        family
        for _, family, marked in family_rows
        if family is not None and marked
    }
    connect_by_key(
        "explicit_derivative_family",
        (
            (sample_id, family if family in derivative_families else None)
            for sample_id, family, _ in family_rows
        ),
    )

    members_by_root: defaultdict[str, list[str]] = defaultdict(list)
    for sample_id in selected:
        members_by_root[union.find(sample_id)].append(sample_id)
    group_by_sample: dict[str, str] = {}
    groups: list[dict[str, object]] = []
    ordered_groups = sorted(
        members_by_root.values(), key=lambda value: tuple(sorted(value))
    )
    for members in ordered_groups:
        members = sorted(members)
        group_id = members[0] if len(members) == 1 else (
            "identity-group-" + sha256_sorted_ids(members)[:16]
        )
        for sample_id in members:
            group_by_sample[sample_id] = group_id
        if len(members) > 1:
            evidence = sorted(
                item
                for (left, right), values in evidence_by_pair.items()
                if left in members and right in members
                for item in values
            )
            groups.append(
                {
                    "group_id": group_id,
                    "sample_ids": members,
                    "sample_ids_sha256": sha256_sorted_ids(members),
                    "evidence": evidence,
                }
            )
    return group_by_sample, groups


def lineage_stratum(record: CCLEIdentityRecord) -> str:
    tissue = _normalise_label(record.tissue_type) or "unknown_tissue"
    haem = _normalise_label(record.haem_lineage)
    if haem not in {None, "na", "n a", "none", "not applicable"}:
        return f"{tissue}|haem:{haem}"
    cancer = _normalise_label(record.cancer_type)
    return f"{tissue}|cancer:{cancer or 'unknown'}"


def _integer_targets(total: int, ratios: Mapping[str, float]) -> dict[str, int]:
    if tuple(ratios) != PARTITION_ORDER:
        raise CCLEContractError(
            f"partition ratio keys/order must be {PARTITION_ORDER!r}"
        )
    if total < 1 or any(value < 0 for value in ratios.values()):
        raise CCLEContractError("invalid split total or ratio")
    if not math.isclose(sum(ratios.values()), 1.0, rel_tol=0.0, abs_tol=1e-12):
        raise CCLEContractError("partition ratios must sum to one")
    raw = {name: total * ratios[name] for name in PARTITION_ORDER}
    result = {name: int(math.floor(raw[name])) for name in PARTITION_ORDER}
    remaining = total - sum(result.values())
    tie_priority = {"validation": 0, "internal_final": 1, "train": 2}
    ranked = sorted(
        PARTITION_ORDER,
        key=lambda name: (-(raw[name] - result[name]), tie_priority[name]),
    )
    for name in ranked[:remaining]:
        result[name] += 1
    return result


def deterministic_group_stratified_split(
    identities: Mapping[str, CCLEIdentityRecord],
    sample_ids: Iterable[str],
    *,
    seed: int = 20260830,
    ratios: Mapping[str, float] = DEFAULT_PARTITION_RATIOS,
) -> tuple[dict[str, tuple[str, ...]], dict[str, object]]:
    """Create a stable tissue/lineage-aware split with identity groups intact."""

    supplied = list(sample_ids)
    selected = sorted(set(supplied))
    if len(selected) != len(supplied):
        raise CCLEContractError("duplicate input sample identifiers")
    group_by_sample, non_singleton_groups = infer_identity_groups(identities, selected)
    members_by_group: defaultdict[str, list[str]] = defaultdict(list)
    for sample_id in selected:
        members_by_group[group_by_sample[sample_id]].append(sample_id)

    groups_by_stratum: defaultdict[
        str, list[tuple[str, tuple[str, ...]]]
    ] = defaultdict(list)
    group_strata: dict[str, str] = {}
    for group_id, members in members_by_group.items():
        stratum_counts = Counter(lineage_stratum(identities[item]) for item in members)
        stratum = sorted(stratum_counts, key=lambda key: (-stratum_counts[key], key))[0]
        group_strata[group_id] = stratum
        groups_by_stratum[stratum].append((group_id, tuple(sorted(members))))

    global_target = _integer_targets(len(selected), ratios)
    global_count = {name: 0 for name in PARTITION_ORDER}
    assignment: dict[str, str] = {}
    for stratum in sorted(groups_by_stratum):
        groups = groups_by_stratum[stratum]
        stratum_total = sum(len(members) for _, members in groups)
        local_target = _integer_targets(stratum_total, ratios)
        local_count = {name: 0 for name in PARTITION_ORDER}
        groups.sort(
            key=lambda item: (
                -len(item[1]),
                hashlib.sha256(
                    f"{seed}|{stratum}|{item[0]}".encode("utf-8")
                ).hexdigest(),
            )
        )
        for group_id, members in groups:
            size = len(members)

            def rank(partition: str) -> tuple[float, str]:
                local_remaining = local_target[partition] - local_count[partition]
                global_remaining = global_target[partition] - global_count[partition]
                local_scale = max(1, local_target[partition])
                global_scale = max(1, global_target[partition])
                fit_bonus = 2.0 if local_remaining >= size else 0.0
                global_fit_bonus = 0.5 if global_remaining >= size else 0.0
                utility = (
                    fit_bonus
                    + global_fit_bonus
                    + local_remaining / local_scale
                    + 0.35 * global_remaining / global_scale
                )
                tie = hashlib.sha256(
                    f"{seed}|{group_id}|{partition}".encode("utf-8")
                ).hexdigest()
                return (-utility, tie)

            partition = min(PARTITION_ORDER, key=rank)
            assignment[group_id] = partition
            local_count[partition] += size
            global_count[partition] += size

    partitions: dict[str, tuple[str, ...]] = {}
    for partition in PARTITION_ORDER:
        partitions[partition] = tuple(
            sorted(
                sample_id
                for sample_id in selected
                if assignment[group_by_sample[sample_id]] == partition
            )
        )
    sets = [set(partitions[name]) for name in PARTITION_ORDER]
    if any(
        sets[left] & sets[right]
        for left in range(3)
        for right in range(left + 1, 3)
    ):
        raise CCLEContractError("split partitions overlap")
    if set().union(*sets) != set(selected):
        raise CCLEContractError("split partitions do not cover the triple universe")
    for group_id, members in members_by_group.items():
        homes = {assignment[group_id] for _ in members}
        if len(homes) != 1:  # pragma: no cover - construction invariant
            raise CCLEContractError(
                f"identity group split across partitions: {group_id}"
            )

    audit = {
        "algorithm": "deterministic_group_stratified_greedy_v1",
        "seed": seed,
        "ratios": dict(ratios),
        "target_counts": global_target,
        "observed_counts": {name: len(partitions[name]) for name in PARTITION_ORDER},
        "stratum_rule": "Tissue_type plus Haem_lineage when present, else Cancer_type",
        "identity_group_rule": (
            "same source Identifier, exact normalized Cell_line, or conservative "
            "explicit parental/resistant/derived/clone family"
        ),
        "identity_group_count": len(members_by_group),
        "non_singleton_identity_groups": non_singleton_groups,
        "strata": {
            stratum: {
                partition: sum(
                    1
                    for sample_id in partitions[partition]
                    if group_strata[group_by_sample[sample_id]] == stratum
                )
                for partition in PARTITION_ORDER
            }
            for stratum in sorted(groups_by_stratum)
        },
    }
    return partitions, audit


def build_ccle_split_manifest(
    inventory: CCLETripleInventory,
    *,
    seed: int = 20260830,
    frozen_on: str = "2026-08-30",
) -> dict[str, object]:
    partitions, split_audit = deterministic_group_stratified_split(
        inventory.identities, inventory.triple_sidms, seed=seed
    )
    partition_payload: dict[str, object] = {}
    for partition, sample_ids in partitions.items():
        partition_payload[partition] = {
            "count": len(sample_ids),
            "sample_ids": list(sample_ids),
            "sample_ids_sha256": sha256_sorted_ids(sample_ids),
            "status": (
                "sealed_untouched_final"
                if partition == "internal_final"
                else "development_visible"
            ),
        }
    manifest: dict[str, object] = {
        "schema": "creed.ccle-l0-l1-l2-split.v1",
        "protocol_id": "CREED_CCLE_L0_L1_L2_SPLIT_V1",
        "status": "prefrozen_unreleased_internal_final",
        "frozen_on": frozen_on,
        "purpose": (
            "Pre-freeze exact CCLE metabolomics-RNA-protein identities before "
            "any CREED model sees assay values."
        ),
        "read_boundary": {
            "performed": [
                "full-file SHA-256",
                "source schemas and headers",
                "sample identifier columns/rows",
                "mmc2 non-expression identity and tissue/lineage metadata",
            ],
            "forbidden_and_not_performed": [
                "metabolite abundance values",
                "RNA expected-count values",
                "protein abundance values",
                "feature selection or outcome inspection",
                "model fitting or model selection",
            ],
        },
        "identity_contract": {
            "sole_authority": (
                "ProCan mmc2 model_id(SIDM) <-> BROAD_ID(ACH) / CCLE_ID bridge"
            ),
            "canonical_sample_id": "mmc2 model_id (SIDM)",
            "metabolomics_resolution": (
                "Resolve both CCLE_ID and DepMap_ID through mmc2; conflicting "
                "resolutions fail closed."
            ),
            "rna_resolution": "source model_id header only",
            "protein_resolution": (
                f"mmc3 {inventory.protein_sheet!r} Project_Identifier SIDM prefix"
            ),
            "forbidden_inputs": [
                "CCLE_DepMap_18Q2_RNAseq_RPKM_20180502.gct",
                "any truncated GCT",
                "any symbol-only or derived-symbol expression matrix",
            ],
            "rna_feature_identity_for_later_value_release": (
                "source gene_id only; symbol is annotation and cannot be the "
                "primary key"
            ),
        },
        "sources": {
            role: asdict(fingerprint)
            for role, fingerprint in sorted(inventory.source_fingerprints.items())
        },
        "sample_universe": {
            "mmc2_identity_samples": len(inventory.identities),
            "mmc3_protein_header_samples": inventory.protein_header_samples,
            "rna_header_samples": inventory.rna_header_samples,
            "metabolomics_rows": inventory.metabolomics_rows,
            "metabolomics_rows_mapped_via_mmc2": len(inventory.metabolomics_sidms),
            "metabolomics_rows_unmapped_via_mmc2": inventory.metabolomics_unmapped_rows,
            "exact_triple_samples": len(inventory.triple_sidms),
            "exact_triple_sample_ids_sha256": sha256_sorted_ids(
                inventory.triple_sidms
            ),
        },
        "split": split_audit,
        "partitions": partition_payload,
        "internal_final_release_gate": {
            "released": False,
            "gbm_role": (
                "CPTAC GBM is external-domain development validation only because "
                "its outcomes were previously exposed during OX development."
            ),
            "must_be_frozen_before_release": [
                "operator definitions and directions",
                "all preprocessing and feature-identity registries",
                "model weights and full asset hashes",
                "training/validation decisions and hyperparameters",
                "direct, equal-cost, oracle, shuffle, and FM-ablation controls",
                "primary/secondary metrics and hierarchical-bootstrap procedure",
                "code commit, environment lock, seeds, and one-shot runner",
            ],
            "release_conditions": [
                "create-once TEMPLAR release receipt binds every item above",
                "no internal-final assay value has been read before that receipt",
                "one evaluation invocation with no automatic retry",
                "all outcomes are retained and reported regardless of direction",
            ],
            "prohibited_before_release": [
                "loading any internal-final metabolite/RNA/protein value",
                "early stopping, thresholding, feature selection, or debugging "
                "on final",
                "changing an operator or analysis after seeing final outcomes",
            ],
        },
        "claim_boundary": {
            "this_artifact_may_support": [
                "identity provenance",
                "sample-disjoint development/validation/final allocation",
                "evidence that the allocation preceded assay-value access",
            ],
            "this_artifact_cannot_support": [
                "scientific performance claims",
                "virtual-layer completion claims",
                "publication or release authorization",
            ],
        },
    }
    manifest["integrity"] = {
        "canonical_payload_sha256": canonical_json_sha256(manifest),
        "canonicalisation": "UTF-8 JSON, sorted keys, compact separators",
    }
    return manifest

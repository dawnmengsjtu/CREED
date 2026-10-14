"""Full identity-safe M+RNA pool extraction (generator-fullpool-20261009-0001), pre-registered in PLAN.md.

Identity bridge: ProCan mmc2 'Cell line level sample info' (authority: SIDM <-> BROAD_ID/CCLE_ID, Tissue_type) extended by DepMap
Model.csv (ModelID <-> SangerModelID <-> PatientID, CCLEName, OncotreeLineage) only where mmc2 has no mapping; mmc2 wins on conflicts;
unmapped lines are excluded, never guessed. Sealed 608 SIDMs (421 train + 64 validation + 123 internal_final) and any line sharing
ACH / PatientID / CCLE name / normalised cell-line name with a 608 line are excluded before any assay value is decoded.
Assay values are decoded column/row-gated: only pool lines and the 5 named 421-train parity lines are ever converted to numbers.
mmc3 is opened only for the sample-ID column (first column) of 'Prot matrix excl single-peptide'; no protein value is stored.

usage: python extract_fullpool.py OUT_DIR
"""
import csv
import gzip
import hashlib
import json
import os
import re
import sys
import time
from collections import Counter, defaultdict

import numpy as np
import openpyxl

RAW = '/public/home/mengxl/creed-research-runs/_raw_whitelist_20261009'
RES = '/public/home/mengxl/creed-research-runs/extra-supervision-20261008-0001/resources'
RNA421 = '/public/home/mengxl/creed-research-runs/rna421-20260906-0002'
SOURCE = '/public/home/mengxl/creed-research-runs/gdsc-metabolite-full-protein-20260908-0001/output/result/source-inputs.npz'
RAW_SHA = {'CCLE_metabolomics_20190502.csv': '7c1d24aa575f4c58a29019026b5df8e6d1142a56925aba32ff3f1d1d5a7fd0ac',
           'CCLE_RNAseq_RSEM_Expected_Counts.csv.gz': 'f59221c04e15a2ba807181d23ba53fa4acd778f02ecb05321c473857e824191f',
           'mmc2.xlsx': '9327f665d259ef4d148d3a64e7c76cdeb51dce8dee06768efd6b7d044130e203',
           'mmc3.xlsx': 'f6c414bc07fdf6e33680cd61c4b3d22d26e531a27eff2cdef3648bfb46e72f34',
           'Model.csv': 'a4cac376131b41aa10b60a075b11c80264bfa860a5509d22cde259c5e85867f8'}
SOURCE_SHA = '52d71517a8952e83a1603cd831f09449544102d0ca5ff829ff01be9bf79cd637'
EXTRA66_SHA = 'bd8d13f57b691f579dfcd6210a44f4ab1e05a8691ced62ff2baeb7a6a67883af'
RNA_CONFIG_SHA = '995a42ab7cbf4d9d9c5db3eb9ecf79d1fe8887bc43841846b9202b4e48ed5a6d'
PARITY = ['SIDM00018', 'SIDM00080', 'SIDM00082', 'SIDM00396', 'SIDM00685']  # 421-train lines only (never 64/123)
GATE = 1e-9
# inner-bank held-out tissues per outer fold, transcribed from generator_scaleup_20261009/INVENTORY.md section 3.2
INNER_HELD = {
    0: {'inner-0': ['Esophagus', 'Haematopoietic and Lymphoid', 'Liver', 'Other tissue', 'Prostate', 'Soft Tissue'],
        'inner-1': ['Biliary Tract', 'Bladder', 'Bone', 'Breast', 'Kidney', 'Pancreas', 'Skin', 'Thyroid'],
        'inner-2': ['Central Nervous System', 'Endometrium', 'Head and Neck', 'Large Intestine', 'Ovary', 'Peripheral Nervous System', 'Stomach']},
    1: {'inner-0': ['Biliary Tract', 'Head and Neck', 'Lung', 'Soft Tissue'],
        'inner-1': ['Bladder', 'Breast', 'Liver', 'Other tissue', 'Pancreas', 'Prostate', 'Skin', 'Stomach'],
        'inner-2': ['Central Nervous System', 'Esophagus', 'Kidney', 'Large Intestine', 'Ovary', 'Peripheral Nervous System', 'Thyroid']},
    2: {'inner-0': ['Biliary Tract', 'Liver', 'Lung', 'Other tissue', 'Soft Tissue'],
        'inner-1': ['Haematopoietic and Lymphoid', 'Ovary', 'Stomach', 'Thyroid'],
        'inner-2': ['Bone', 'Central Nervous System', 'Endometrium', 'Esophagus', 'Head and Neck', 'Large Intestine', 'Skin']},
    3: {'inner-0': ['Liver', 'Lung', 'Other tissue', 'Peripheral Nervous System'],
        'inner-1': ['Bladder', 'Bone', 'Endometrium', 'Haematopoietic and Lymphoid', 'Pancreas'],
        'inner-2': ['Breast', 'Central Nervous System', 'Esophagus', 'Kidney', 'Prostate', 'Skin', 'Thyroid']},
    4: {'inner-0': ['Biliary Tract', 'Head and Neck', 'Lung', 'Soft Tissue'],
        'inner-1': ['Bladder', 'Bone', 'Endometrium', 'Haematopoietic and Lymphoid', 'Pancreas'],
        'inner-2': ['Breast', 'Kidney', 'Large Intestine', 'Ovary', 'Peripheral Nervous System', 'Prostate', 'Stomach']}}


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for b in iter(lambda: fh.read(1 << 20), b''):
            h.update(b)
    return h.hexdigest()


def prefix_fields(raw, last):
    """Quote-aware byte split of the first `last`+1 CSV fields (verbatim from extract_m_gdsc.py)."""
    out = []; start = 0; i = 0; quoted = False
    while i < len(raw):
        c = raw[i]
        if c == 34:
            if quoted and i + 1 < len(raw) and raw[i + 1] == 34:
                i += 2; continue
            quoted = not quoted
        if c == 44 and not quoted:
            out.append(raw[start:i]); start = i + 1
            if len(out) > last:
                return out
        i += 1
    out.append(raw[start:].rstrip(b'\r\n')); assert len(out) > last and not quoted
    return out


def text(v):
    return next(csv.reader([v.decode('utf-8-sig')]))[0].strip()


def number(v):
    if v.strip().lower() in ('', 'na', 'nan', 'null', 'none'):
        return np.nan
    return float(v)


def norm_name(s):
    return re.sub(r'[^A-Z0-9]', '', (s or '').upper())


def log(**kw):
    print(json.dumps(kw, ensure_ascii=False), flush=True)


def main(out):
    t0 = time.time()
    os.makedirs(out + '/pack', exist_ok=False)
    # ---------------------------------------------------------------- 0. fingerprints
    for name, h in RAW_SHA.items():
        assert sha(f'{RAW}/{name}') == h, name
    assert sha(SOURCE) == SOURCE_SHA and sha(f'{RES}/m_rna_extra66.npz') == EXTRA66_SHA and sha(f'{RES}/rna_config.json') == RNA_CONFIG_SHA
    log(event='fingerprints_ok')
    split = json.load(open(os.path.dirname(os.path.abspath(__file__)) + '/../resources/original_identity_split_ids.json'))
    train, val, fin = (split[k] for k in ('train', 'validation', 'internal_final'))
    cfg = json.load(open(f'{RES}/rna_config.json'))
    old608 = set(cfg['excluded_old608'])
    assert len(train) == 421 and len(val) == 64 and len(fin) == 123
    assert set(train) | set(val) | set(fin) == old608 and len(old608) == 608
    assert set(train) <= old608 and not (set(train) & set(val)) and not (set(train) & set(fin)) and not (set(val) & set(fin))
    sealed = set(val) | set(fin)
    assert set(PARITY) <= set(train) and not (set(PARITY) & sealed)
    S = np.load(SOURCE, allow_pickle=False)
    assert sorted(S['sample_ids'].tolist()) == sorted(train)
    rna421_ids = np.load(f'{RNA421}/sample_ids.npy', allow_pickle=False)
    assert sorted(rna421_ids.tolist()) == sorted(train)
    gene_ids = np.load(f'{RNA421}/gene_ids.npy', allow_pickle=False)
    all_gene_ids = np.load(f'{RNA421}/all_source_gene_ids.npy', allow_pickle=False)
    common = np.load(f'{RNA421}/common_source_gene_indices.npy', allow_pickle=False)
    assert len(common) == 34985 and len(set(common.tolist())) == 34985 and (all_gene_ids[common] == gene_ids).all()
    assert common.tolist() == cfg['common_source_gene_indices']
    log(event='sealed_sets_ok', train=len(train), validation=len(val), internal_final=len(fin), union=len(old608))

    # ---------------------------------------------------------------- 1. identity tables (IDs / metadata only)
    wb = openpyxl.load_workbook(f'{RAW}/mmc2.xlsx', read_only=True)
    ws = wb['Cell line level sample info']
    rows = ws.iter_rows(values_only=True)
    next(rows); hdr = [str(c) if c is not None else '' for c in next(rows)]
    col = {k: hdr.index(k) for k in ('model_id', 'Cell_line', 'Source', 'Identifier', 'Tissue_type', 'BROAD_ID', 'CCLE_ID')}
    mmc2 = {}
    for r in rows:
        if r is None or r[col['model_id']] is None:
            continue
        sid = str(r[col['model_id']]).strip()
        assert re.fullmatch(r'SIDM\d{5}', sid) and sid not in mmc2, sid
        g = lambda k: ('' if r[col[k]] is None else str(r[col[k]]).strip())
        mmc2[sid] = dict(cell_line=g('Cell_line'), source=g('Source'), identifier=g('Identifier'), tissue=g('Tissue_type'), ach=g('BROAD_ID'), ccle=g('CCLE_ID'))
    wb.close()
    assert len(mmc2) == 949, len(mmc2)
    mmc2_ach = defaultdict(set); mmc2_ccle = defaultdict(set); multi_ach = {}; multi_ccle = {}
    for sid, r in mmc2.items():
        r['achs'] = [a.strip() for a in r['ach'].split(';') if a.strip()]; r['ccles'] = [c.strip() for c in r['ccle'].split(';') if c.strip()]
        for a in r['achs']:
            assert re.fullmatch(r'ACH-\d{6}', a), r['ach']; mmc2_ach[a].add(sid)
        for c in r['ccles']:
            mmc2_ccle[c].add(sid)
        if len(r['achs']) > 1:
            multi_ach[sid] = r['achs']
        if len(r['ccles']) > 1:
            multi_ccle[sid] = r['ccles']
    mmc2_ach_ambig = {a: sorted(s) for a, s in mmc2_ach.items() if len(s) > 1}
    mmc2_ccle_ambig = {c: sorted(s) for c, s in mmc2_ccle.items() if len(s) > 1}
    mmc2_ach2sid = {a: next(iter(s)) for a, s in mmc2_ach.items() if len(s) == 1}
    mmc2_ccle2sid = {c: next(iter(s)) for c, s in mmc2_ccle.items() if len(s) == 1}
    assert set(old608) <= set(mmc2)
    log(event='mmc2_ok', identities=len(mmc2), with_ach=sum(1 for r in mmc2.values() if r['ach']), with_ccle=sum(1 for r in mmc2.values() if r['ccle']),
        ach_ambiguous=mmc2_ach_ambig, ccle_ambiguous=mmc2_ccle_ambig, multi_valued_broad_id=multi_ach, multi_valued_ccle_id=multi_ccle, tissue_vocab=sorted(set(r['tissue'] for r in mmc2.values())))

    model = {}
    with open(f'{RAW}/Model.csv', encoding='utf-8-sig', newline='') as fh:
        rd = csv.DictReader(fh)
        for r in rd:
            a = r['ModelID'].strip(); assert re.fullmatch(r'ACH-\d{6}', a) and a not in model, a
            model[a] = dict(patient=r['PatientID'].strip(), name=r['CellLineName'].strip(), stripped=r['StrippedCellLineName'].strip(),
                            lineage=r['OncotreeLineage'].strip(), ccle=r['CCLEName'].strip(), sidm=r['SangerModelID'].strip())
    assert len(model) == 1959, len(model)
    model_sid = defaultdict(set); model_ccle = defaultdict(set)
    for a, r in model.items():
        if r['sidm']:
            assert re.fullmatch(r'SIDM\d{5}', r['sidm']), r['sidm']; model_sid[r['sidm']].add(a)
        if r['ccle']:
            model_ccle[r['ccle']].add(a)
    model_sid_ambig = {s: sorted(v) for s, v in model_sid.items() if len(v) > 1}
    model_ccle_ambig = {c: sorted(v) for c, v in model_ccle.items() if len(v) > 1}
    model_sid2ach = {s: next(iter(v)) for s, v in model_sid.items() if len(v) == 1}
    model_ccle2ach = {c: next(iter(v)) for c, v in model_ccle.items() if len(v) == 1}
    log(event='model_csv_ok', models=len(model), with_sidm=sum(1 for r in model.values() if r['sidm']), sidm_ambiguous=model_sid_ambig, ccle_ambiguous=model_ccle_ambig)

    # bridge SIDM <-> ACH: mmc2 first; Model.csv only where mmc2 is silent and nothing conflicts
    sid2ach = {}; ach2sid = {}; bridge_src = {}; conflicts = []; sid_achs = defaultdict(list)
    for a, s in mmc2_ach2sid.items():
        ach2sid[a] = s; bridge_src[s] = 'mmc2'; sid_achs[s].append(a)
    for s in sid_achs:  # first listed unambiguous ACH is the representative; all listed ACHs stay in sid_achs for exclusion keys
        sid_achs[s] = [a for a in mmc2[s]['achs'] if a in sid_achs[s]]; sid2ach[s] = sid_achs[s][0]
    for s, a in model_sid2ach.items():
        if s in sid2ach:
            if sid2ach[s] != a:
                conflicts.append(dict(sidm=s, mmc2_ach=sid2ach[s], model_csv_ach=a, resolution='mmc2 wins'))
            continue
        if a in ach2sid:  # mmc2 maps this ACH to another SIDM
            conflicts.append(dict(sidm=s, model_csv_ach=a, mmc2_sidm_for_that_ach=ach2sid[a], resolution='mmc2 wins; Model.csv SIDM unmapped'))
            continue
        if a in mmc2_ach_ambig:  # mmc2 lists this ACH under several SIDMs: fail closed
            conflicts.append(dict(sidm=s, model_csv_ach=a, mmc2_sidms_for_that_ach=mmc2_ach_ambig[a], resolution='mmc2 ambiguous; unmapped'))
            continue
        sid2ach[s] = a; ach2sid[a] = s; bridge_src[s] = 'model_csv'; sid_achs[s] = [a]
    # mmc2 lines with a SIDM but no BROAD_ID whose Model.csv SIDM points to an ACH: covered above (model_csv source).
    log(event='bridge_ok', pairs=len(sid2ach), from_mmc2=sum(v == 'mmc2' for v in bridge_src.values()), from_model_csv=sum(v == 'model_csv' for v in bridge_src.values()), conflicts=conflicts)

    def patient(s):
        a = sid2ach.get(s); return model[a]['patient'] if a in model else ''

    def patients_all(s):
        return {model[a]['patient'] for a in sid_achs.get(s, []) if a in model and model[a]['patient']}

    def lineage(s):
        a = sid2ach.get(s); return model[a]['lineage'] if a in model else ''

    def ccle_names(s):
        names = set()
        if s in mmc2:
            names |= set(mmc2[s]['ccles'])
        for a in sid_achs.get(s, []):
            if a in model and model[a]['ccle']:
                names.add(model[a]['ccle'])
        return names

    def norm_names(s):
        names = set()
        if s in mmc2 and mmc2[s]['cell_line']:
            names.add(norm_name(mmc2[s]['cell_line']))
        for a in sid_achs.get(s, []):
            if a in model:
                names |= {norm_name(model[a]['name']), norm_name(model[a]['stripped'])}
        names.discard('')
        return names

    # lineage -> project tissue vocabulary, learned from lines present in both tables (majority, >=90% agreement)
    vocab = set(S['tissues'].tolist()); assert vocab <= set(r['tissue'] for r in mmc2.values())
    lin2tis = defaultdict(Counter)
    for s, r in mmc2.items():
        a = sid2ach.get(s)
        if a in model and model[a]['lineage'] and r['tissue']:
            lin2tis[model[a]['lineage']][r['tissue']] += 1
    lineage_map = {}; lineage_map_detail = {}
    for lin, c in lin2tis.items():
        tis, n = c.most_common(1)[0]; frac = n / sum(c.values())
        lineage_map_detail[lin] = dict(majority=tis, n=n, total=sum(c.values()), frac=round(frac, 3), counts=dict(c))
        if frac >= .9 and sum(c.values()) >= 3:
            lineage_map[lin] = tis

    def tissue_of(s):
        if s in mmc2 and mmc2[s]['tissue']:
            return mmc2[s]['tissue'], 'mmc2'
        lin = lineage(s)
        if lin in lineage_map:
            return lineage_map[lin], 'model_csv_lineage'
        return '', 'unresolved'

    # ---------------------------------------------------------------- 2. sealed exclusion keys (IDs only; never written out)
    old_ach = set().union(*(set(sid_achs.get(s, [])) for s in old608))
    old_pat = set().union(*(patients_all(s) for s in old608))
    old_ccle = set().union(*(ccle_names(s) for s in old608))
    old_norm = set().union(*(norm_names(s) for s in old608))
    log(event='sealed_keys', ach=len(old_ach), patient=len(old_pat), ccle_names=len(old_ccle), norm_names=len(old_norm))

    # ---------------------------------------------------------------- 3. metabolomics identity pass (no values decoded)
    mp = f'{RAW}/CCLE_metabolomics_20190502.csv'
    with open(mp, 'rb') as f:
        m_header = next(csv.reader([f.readline().decode('utf-8-sig')]))
        assert m_header[:2] == ['CCLE_ID', 'DepMap_ID'] and m_header[2:] == S['metabolite_ids'].tolist() == cfg['m_feature_ids']
        m_rows = []
        for raw in f:
            if not raw.strip():
                continue
            p = prefix_fields(raw, 1); m_rows.append((text(p[0]), text(p[1])))
    assert len(m_rows) == 928
    res_count = Counter(); m_res = []  # per source row: resolved SIDM via bridge, mmc2-only SIDM, and reason
    for i, (ccle, dep) in enumerate(m_rows):
        # mmc2-only resolution (protocol of the 608): both ids through mmc2, conflict fails closed
        s_dep2 = mmc2_ach2sid.get(dep) if dep else None
        s_ccle2 = mmc2_ccle2sid.get(ccle) if ccle else None
        if s_dep2 and s_ccle2 and s_dep2 != s_ccle2:
            s_mmc2 = None; r2 = 'mmc2_conflict'
        else:
            s_mmc2 = s_dep2 or s_ccle2; r2 = 'mmc2' if s_mmc2 else 'mmc2_unmapped'
        # bridge resolution: DepMap_ID -> bridge; CCLE_ID -> mmc2 CCLE_ID or Model.csv CCLEName -> ACH -> bridge
        s_dep = ach2sid.get(dep) if dep else None
        s_ccle = s_ccle2
        if s_ccle is None and ccle in model_ccle2ach:
            s_ccle = ach2sid.get(model_ccle2ach[ccle])
        if s_dep and s_ccle and s_dep != s_ccle:
            s = None; reason = 'bridge_conflict'
        elif s_dep or s_ccle:
            s = s_dep or s_ccle; reason = 'both' if (s_dep and s_ccle) else ('depmap_only' if s_dep else 'ccle_only')
        else:
            s = None; reason = 'unmapped'
        res_count[reason] += 1
        m_res.append(dict(row=i, ccle=ccle, dep=dep, sidm=s, reason=reason, sidm_mmc2=s_mmc2, reason_mmc2=r2))
    m_sidms_mmc2 = {r['sidm_mmc2'] for r in m_res if r['sidm_mmc2']}
    m_sidms_bridge = {r['sidm'] for r in m_res if r['sidm']}
    log(event='m_identity', rows=928, bridge_resolution=dict(res_count), mmc2_resolution=dict(Counter(r['reason_mmc2'] for r in m_res)),
        unique_sidm_bridge=len(m_sidms_bridge), unique_sidm_mmc2=len(m_sidms_mmc2))
    assert old608 <= m_sidms_mmc2
    # whole-string convention (multi-valued BROAD_ID/CCLE_ID fields never match a single M id): the 608 protocol's 609 figure
    whole_ach = {r['ach']: s for s, r in mmc2.items() if r['ach'] and ';' not in r['ach']}
    whole_ccle = {r['ccle']: s for s, r in mmc2.items() if r['ccle'] and ';' not in r['ccle']}
    m_sidms_mmc2_whole = set()
    for ccle, dep in m_rows:
        a = whole_ach.get(dep) if dep else None; b = whole_ccle.get(ccle) if ccle else None
        if a and b and a != b:
            continue
        if a or b:
            m_sidms_mmc2_whole.add(a or b)
    assert old608 <= m_sidms_mmc2_whole
    mmc2_only_outside = sorted(m_sidms_mmc2 - old608); mmc2_whole_outside = sorted(m_sidms_mmc2_whole - old608)
    log(event='m_mmc2_only', split_convention=len(m_sidms_mmc2), split_outside_608=mmc2_only_outside, whole_string_convention=len(m_sidms_mmc2_whole),
        whole_outside_608=mmc2_whole_outside, protocol_figure=609)

    # ---------------------------------------------------------------- 4. RNA header + protein sample ids (identity only)
    rp = f'{RAW}/CCLE_RNAseq_RSEM_Expected_Counts.csv.gz'
    with gzip.open(rp, 'rb') as f:
        head = [next(csv.reader([f.readline().decode('utf-8-sig')])) for _ in range(4)]
    assert [h[0] for h in head] == ['model_id', 'model_name', 'dataset_name', 'gene_id'] and head[0][1] == ''
    rna_cols = {s: j for j, s in enumerate(head[0][2:], 2)}; rna_batch = {s: head[2][j] for s, j in rna_cols.items()}
    assert len(rna_cols) == 1047 == len(head[0]) - 2 and all(re.fullmatch(r'SIDM\d{5}', s) for s in rna_cols)
    assert old608 <= set(rna_cols)
    wb3 = openpyxl.load_workbook(f'{RAW}/mmc3.xlsx', read_only=True)
    ws3 = wb3['Prot matrix excl single-peptide']
    p_sidms = []; p_first = []
    for k, r in enumerate(ws3.iter_rows(min_col=1, max_col=1, values_only=True)):
        v = '' if r[0] is None else str(r[0]).strip()
        if k < 3:
            p_first.append(v[:40])
        m_ = re.match(r'(SIDM\d{5})', v)
        if m_:
            p_sidms.append(m_.group(1))
    wb3.close()
    assert len(p_sidms) == len(set(p_sidms)) == 949, len(p_sidms)
    p_set = set(p_sidms); assert old608 <= p_set
    log(event='rna_protein_headers_ok', rna_header=len(rna_cols), protein_rows=len(p_sidms), protein_first_cells=p_first, t=round(time.time() - t0, 1))

    # ---------------------------------------------------------------- 5. candidate pool
    cand = {}  # sidm -> record (first source row kept)
    dup_rows = []
    for r in m_res:
        s = r['sidm']
        if s is None or s not in rna_cols:
            continue
        if s in cand:
            dup_rows.append(dict(sidm=s, kept_row=cand[s]['m_row'], dropped_row=r['row'], ccle=r['ccle'], dep=r['dep']))
            continue
        cand[s] = dict(sidm=s, m_row=r['row'], m_ccle=r['ccle'], m_dep=r['dep'], resolution=r['reason'], ach=(r['dep'] if r['dep'] in sid_achs[s] else sid2ach[s]), bridge=bridge_src[s])
    n_cand = len(cand)
    excl = {}; order = ['sealed_608', 'shares_ach_with_608', 'shares_patient_with_608', 'shares_ccle_name_with_608', 'shares_cell_name_with_608']
    for s, c in cand.items():
        a = c['ach']; pt = patient(s)
        if s in old608:
            why = 'sealed_608'
        elif set(sid_achs.get(s, [])) & old_ach:
            why = 'shares_ach_with_608'
        elif patients_all(s) & old_pat:
            why = 'shares_patient_with_608'
        elif ccle_names(s) & old_ccle:
            why = 'shares_ccle_name_with_608'
        elif norm_names(s) & old_norm:
            why = 'shares_cell_name_with_608'
        else:
            continue
        excl[s] = why
    pool = [c for s, c in sorted(cand.items()) if s not in excl]
    # within-pool duplicate guard: one per ACH, one per PatientID (keep lowest SIDM)
    seen_a, seen_p, pool2, dup_pool = set(), set(), [], []
    for c in pool:
        a = c['ach']; pt = patient(c['sidm'])
        if a in seen_a:
            dup_pool.append(dict(sidm=c['sidm'], key='ach', value=a)); continue
        if pt and pt in seen_p:
            dup_pool.append(dict(sidm=c['sidm'], key='patient', value=pt)); continue
        seen_a.add(a); seen_p.add(pt); pool2.append(c)
    pool = pool2
    for c in pool:
        c['patient'] = patient(c['sidm']); c['lineage'] = lineage(c['sidm']); c['tissue'], c['tissue_src'] = tissue_of(c['sidm'])
        c['rna_batch'] = rna_batch[c['sidm']]; c['in_mmc2'] = c['sidm'] in mmc2; c['has_p_header'] = c['sidm'] in p_set
    ids = [c['sidm'] for c in pool]; assert len(ids) == len(set(ids)) and not (set(ids) & old608) and not (set(ids) & sealed)
    ex_count = Counter(excl.values())
    E = np.load(f'{RES}/m_rna_extra66.npz', allow_pickle=False)
    e66 = E['sample_ids'].tolist(); assert sorted(e66) == sorted(cfg['m_rna_ids']) and len(e66) == 66
    missing66 = sorted(set(e66) - set(ids)); extra_vs66 = sorted(set(ids) - set(e66))
    e66_t = dict(zip(e66, E['tissues'].tolist())); e66_m = dict(zip(e66, E['model_ids'].tolist())); e66_p = dict(zip(e66, E['patient_ids'].tolist()))
    tissue_diff = [dict(sidm=c['sidm'], pool=c['tissue'], pack66=e66_t[c['sidm']], src=c['tissue_src'], lineage=c['lineage']) for c in pool if c['sidm'] in e66_t and e66_t[c['sidm']] != c['tissue']]
    id_diff = [dict(sidm=c['sidm'], pool_ach=c['ach'], pack66_ach=e66_m[c['sidm']], pool_patient=c['patient'], pack66_patient=e66_p[c['sidm']]) for c in pool
               if c['sidm'] in e66_m and (e66_m[c['sidm']] != c['ach'] or e66_p[c['sidm']] != c['patient'])]
    excluded66 = {s: excl.get(s, 'not_candidate') for s in missing66}
    log(event='pool_identity', candidates=n_cand, excluded=dict(ex_count), duplicates_source_rows=len(dup_rows), duplicates_in_pool=dup_pool, pool=len(pool),
        blank_tissue=sum(c['tissue'] == '' for c in pool), audited66_missing=excluded66, pool_not_in_66=extra_vs66, tissue_diff_vs_66=tissue_diff, id_diff_vs_66=id_diff,
        bridge_src=dict(Counter(c['bridge'] for c in pool)), resolution=dict(Counter(c['resolution'] for c in pool)), tissue_src=dict(Counter(c['tissue_src'] for c in pool)))

    # ---------------------------------------------------------------- 6. values (row/column gated) for pool + parity
    allids = ids + PARITY; assert len(set(allids)) == len(allids)
    idx = {s: i for i, s in enumerate(allids)}
    row2idx = {}
    for c in pool:
        row2idx[c['m_row']] = idx[c['sidm']]
    for r in m_res:  # parity rows: 421-train lines, resolved through mmc2
        if r['sidm'] in PARITY and r['row'] not in row2idx:
            assert r['sidm_mmc2'] == r['sidm']; row2idx[r['row']] = idx[r['sidm']]
    m = np.full((len(allids), 225), np.nan); seen = set()
    with open(mp, 'rb') as f:
        f.readline()
        i = -1
        for raw in f:
            if not raw.strip():
                continue
            i += 1
            k = row2idx.get(i)
            if k is None:
                continue  # identity-only traversal; assay fields of every other line stay undecoded
            assert k not in seen; v = next(csv.reader([raw.decode('utf-8-sig')])); assert len(v) == 227
            m[k] = [number(z) for z in v[2:]]; seen.add(k)
    assert len(seen) == len(allids) and not np.isinf(m).any()
    cols = [rna_cols[s] for s in allids]; width = len(head[0])
    counts = np.full((len(allids), len(common)), np.nan); missing_full = np.zeros(len(allids), int); genes = []; slot = {int(g): j for j, g in enumerate(common)}
    with gzip.open(rp, 'rb') as f:
        for _ in range(4):
            f.readline()
        for raw in f:
            if not raw.strip():
                continue
            flds = raw.rstrip(b'\r\n').split(b',')
            if len(flds) != width:
                flds = prefix_fields(raw, width - 1)
            assert len(flds) == width
            gi = len(genes); genes.append(text(flds[0]) if flds[0].strip() else '')
            vals = np.array([number(text(flds[j]) if flds[j].strip() else '') for j in cols], float)
            assert not np.isinf(vals).any() and not np.any(vals < 0)
            missing_full += np.isnan(vals)
            if gi in slot:
                counts[:, slot[gi]] = vals
    assert len(genes) == 37279 and np.array_equal(np.array(genes), all_gene_ids) and np.array_equal(np.array(genes)[common], gene_ids)
    good = np.isfinite(counts).all(1) & (np.nansum(counts, 1) > 0)
    rna = np.full_like(counts, np.nan); rna[good] = np.log1p(counts[good] / counts[good].sum(1)[:, None] * 1e6)
    log(event='values_decoded', rows=len(allids), m_missing=int(np.isnan(m[:len(ids)]).sum()), rna_common_missing=int(np.isnan(counts[:len(ids)]).sum()),
        rna_valid=int(good[:len(ids)].sum()), t=round(time.time() - t0, 1))

    # ---------------------------------------------------------------- 7. hard gates
    si = {s: i for i, s in enumerate(S['sample_ids'].tolist())}; ri = {s: i for i, s in enumerate(rna421_ids.tolist())}
    R421 = np.load(f'{RNA421}/rna_log_cpm.npy', allow_pickle=False, mmap_mode='r')
    gate_m = max(float(np.nanmax(np.abs(m[idx[s]] - S['m'][si[s]]))) for s in PARITY)
    gate_m_nan = all(np.array_equal(np.isnan(m[idx[s]]), np.isnan(S['m'][si[s]])) for s in PARITY)
    gate_r = max(float(np.max(np.abs(rna[idx[s]] - np.asarray(R421[ri[s]])))) for s in PARITY)
    gates = dict(parity_ids=PARITY, m_max_abs_diff=gate_m, m_nan_pattern_equal=gate_m_nan, rna_max_abs_diff=gate_r, threshold=GATE,
                 gene_panel_equal_rna421=True, metabolite_order_equal_source_inputs=True, passed=bool(gate_m <= GATE and gate_m_nan and gate_r <= GATE))
    log(event='gates', **gates)
    assert gates['passed'], gates
    # extra66 numerical cross-check (informational, same transformation expected)
    e_idx = {s: i for i, s in enumerate(e66)}; common66 = [s for s in ids if s in e_idx]
    d66_m = max(float(np.nanmax(np.abs(m[idx[s]] - E['m'][e_idx[s]]))) for s in common66) if common66 else None
    d66_r = max(float(np.max(np.abs(rna[idx[s]] - E['rna_log_cpm'][e_idx[s]]))) for s in common66) if common66 else None
    log(event='extra66_crosscheck', n=len(common66), m_max_abs_diff=d66_m, rna_max_abs_diff=d66_r)

    # ---------------------------------------------------------------- 8. per-fold availability (held-out-tissue scenario)
    tis = np.array([c['tissue'] for c in pool]); fold_t = {int(f): sorted(set(S['tissues'][S['fold_ids'] == f].tolist())) for f in np.unique(S['fold_ids'])}
    folds = {}
    for f in range(5):
        held = set(fold_t[f]); outer = int(np.sum([(t != '') and (t not in held) for t in tis]))
        inner = {b: int(np.sum([(t != '') and (t not in set(INNER_HELD[f][b])) for t in tis])) for b in INNER_HELD[f]}
        folds[f'fold-{f}'] = dict(held_out_tissues=fold_t[f], new_cell_usable=len(pool), tissue_outer_usable=outer, tissue_inner_usable=inner,
                                  excluded_held_tissue=int(np.sum([t in held for t in tis])), excluded_blank=int(np.sum(tis == '')))
    log(event='folds', **{k: v['tissue_outer_usable'] for k, v in folds.items()})

    # ---------------------------------------------------------------- 9. M∩P universe (identity only)
    rna_set = set(rna_cols)
    mp_mmc2 = m_sidms_mmc2 & p_set; mp_bridge = m_sidms_bridge & p_set
    mrp_mmc2 = mp_mmc2 & rna_set; mrp_bridge = mp_bridge & rna_set
    mr_bridge = m_sidms_bridge & rna_set; mr_mmc2 = m_sidms_mmc2 & rna_set
    def outside(S_):
        o = sorted(S_ - old608); return dict(n=len(S_), outside_608=len(o), outside_ids=o if len(o) <= 20 else o[:20] + ['...'])
    universe = dict(mmc2_only=dict(M=len(m_sidms_mmc2), M_and_P=outside(mp_mmc2), M_and_RNA=outside(mr_mmc2), M_RNA_P=outside(mrp_mmc2)),
                    bridge=dict(M=len(m_sidms_bridge), M_and_P=outside(mp_bridge), M_and_RNA=outside(mr_bridge), M_RNA_P=outside(mrp_bridge)),
                    P_header=len(p_set), RNA_header=len(rna_set), note='identity sets before patient/name exclusions; P from mmc3 sample-ID column only')
    for key in ('M_and_P', 'M_RNA_P'):
        for u in ('mmc2_only', 'bridge'):
            o = set(universe[u][key]['outside_ids']) - {'...'}
            universe[u][key]['outside_excluded_reasons'] = {s: excl.get(s, 'not_rna_candidate' if s not in cand else 'in_pool') for s in o}
    log(event='universe', **{k: v for k, v in universe.items() if k != 'note'})

    # ---------------------------------------------------------------- 10. save
    n = len(ids)
    qa = dict(status='extracted', run='generator-fullpool-20261009-0001', script_sha256=sha(os.path.abspath(__file__)), raw_sha256=RAW_SHA, source_inputs_sha256=SOURCE_SHA,
              extra66_sha256=EXTRA66_SHA, rna_config_sha256=RNA_CONFIG_SHA, sealed=dict(train=421, validation=64, internal_final=123, union=608, train_subset_of_608=True, sealed_values_decoded=0),
              identity=dict(mmc2_identities=len(mmc2), mmc2_ach_ambiguous=mmc2_ach_ambig, mmc2_ccle_ambiguous=mmc2_ccle_ambig, model_csv_rows=len(model), model_sidm_ambiguous=model_sid_ambig,
                            bridge_pairs=len(sid2ach), bridge_from_mmc2=sum(v == 'mmc2' for v in bridge_src.values()), bridge_from_model_csv=sum(v == 'model_csv' for v in bridge_src.values()),
                            bridge_conflicts=conflicts, lineage_to_tissue_map=lineage_map, lineage_map_detail=lineage_map_detail),
              metabolomics=dict(rows=928, bridge_resolution=dict(res_count), mmc2_resolution=dict(Counter(r['reason_mmc2'] for r in m_res)), unique_sidm_bridge=len(m_sidms_bridge), unique_sidm_mmc2=len(m_sidms_mmc2),
                                duplicate_source_rows=dup_rows),
              m_mmc2_only=dict(split_convention=len(m_sidms_mmc2), split_outside_608=mmc2_only_outside, whole_string_convention=len(m_sidms_mmc2_whole), whole_outside_608=mmc2_whole_outside, protocol_figure=609),
              rna_header=len(rna_cols), protein_header=len(p_sidms),
              pool=dict(candidates_with_rna=n_cand, excluded=dict(ex_count), excluded_ids={s: w for s, w in excl.items() if s not in old608}, excluded_608_count=sum(1 for w in excl.values() if w == 'sealed_608'),
                        duplicates_in_pool=dup_pool, n=n, blank_tissue=int((tis == '').sum()), blank_tissue_ids=[c['sidm'] for c in pool if c['tissue'] == ''],
                        tissues=dict(Counter(tis.tolist())), rna_batches=dict(Counter(c['rna_batch'] for c in pool)), bridge_src=dict(Counter(c['bridge'] for c in pool)),
                        resolution=dict(Counter(c['resolution'] for c in pool)), tissue_src=dict(Counter(c['tissue_src'] for c in pool)), in_mmc2=sum(c['in_mmc2'] for c in pool), has_p_header=sum(c['has_p_header'] for c in pool),
                        m_missing=int(np.isnan(m[:n]).sum()), rna_common_missing=int(np.isnan(counts[:n]).sum()), rna_full_source_missing=missing_full[:n].tolist(), rna_valid_original_transform=int(good[:n].sum())),
              audited66=dict(subset_of_pool=not missing66, missing=excluded66, pool_not_in_66=extra_vs66, tissue_diff=tissue_diff, id_diff=id_diff, m_max_abs_diff=d66_m, rna_max_abs_diff=d66_r),
              gates=gates, folds=folds, universe=universe, usable_new_cell=n, usable_tissue_per_fold={k: v['tissue_outer_usable'] for k, v in folds.items()},
              feasibility=dict(threshold=150, usable_new_cell=n, proceed=bool(n >= 150)),
              transformation=dict(m='identity: source values preserved (float64); no log, no normalisation, no imputation',
                                  rna='ln(1 + counts / sample_total_of_frozen_34985_common_assayed_genes * 1e6); sample-local; stored float32'),
              sample_table=[dict(sample_id=c['sidm'], ach_id=c['ach'], patient_id=c['patient'], ccle_id=c['m_ccle'], tissue=c['tissue'], tissue_src=c['tissue_src'], lineage=c['lineage'], bridge=c['bridge'],
                                 resolution=c['resolution'], rna_batch=c['rna_batch'], in_mmc2=c['in_mmc2'], has_p_header=c['has_p_header'], in_audited66=c['sidm'] in e_idx) for c in pool],
              seconds=round(time.time() - t0, 1))
    np.savez_compressed(out + '/pack/m_rna_fullpool.npz', sample_ids=np.array(ids), ach_ids=np.array([c['ach'] for c in pool]), patient_ids=np.array([c['patient'] for c in pool]),
                        tissues=np.array(tis, dtype=str), lineages=np.array([c['lineage'] for c in pool]), ccle_ids=np.array([c['m_ccle'] for c in pool]), rna_source_names=np.array([c['rna_batch'] for c in pool]),
                        m=m[:n], rna=rna[:n].astype(np.float32), rna_counts_common=counts[:n].astype(np.float32), metabolite_ids=S['metabolite_ids'], gene_ids=gene_ids, common_source_gene_indices=common,
                        parity_sample_ids=np.array(PARITY), parity_m=m[n:], parity_rna=rna[n:].astype(np.float32), qa_json=np.array(json.dumps(qa)))
    json.dump(qa, open(out + '/pack/qa.json', 'w'), indent=1, ensure_ascii=False)
    log(event='saved', pack=out + '/pack/m_rna_fullpool.npz', n=n, seconds=qa['seconds'])


if __name__ == '__main__':
    main(sys.argv[1])
